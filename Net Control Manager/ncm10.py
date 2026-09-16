import sys
import csv
import json
import os
import re
import threading
import urllib.parse
import urllib.request
from datetime import datetime, timezone
import openpyxl

from PyQt6.QtCore import Qt, QTimer, QAbstractItemModel, QModelIndex
from PyQt6.QtGui import QFont, QTextDocument, QColor, QKeySequence, QShortcut, QFontMetrics
from PyQt6.QtPrintSupport import QPrinter
from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QPushButton, QLineEdit, QComboBox, QTableWidget,
    QTableWidgetItem, QTabWidget, QTextEdit, QGroupBox, QSplitter,
    QFileDialog, QMessageBox, QHeaderView, QFrame, QCheckBox, QDialog,
    QListWidget, QListWidgetItem, QInputDialog, QLayout, QCompleter,
    QSizePolicy, QGridLayout, QScrollArea
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

PROFILES_FILE = os.path.join(BASE_DIR, "profiles.json")
STATION_MEMORY_FILE = os.path.join(BASE_DIR, "station_memory.json")

DEFAULT_PROFILES = {
    "Standard ARES Net": {
        "mode": "Regular",
        "has_roll_call": True,
        "use_custom_statuses": False,
        "use_24hr_clock": False,
        "repeater": {
            "freq": "146.700",
            "offset": "-",
            "tone": "88.5 PL"
        },
        "auth": {
            "enabled": False,
            "users": [
                {"callsign": "", "passcode": "1234"}
            ]
        },
        "station_types": ["Base", "Mobile", "NCS", "Alternate NCS", "NTS", "OES"],
        "traffic_types": ["Question", "Announcement", "Priority", "Emergency", "General Note"],
        "web_status_options": ["ACTIVE", "STANDBY", "CLOSED", "-Normal Operations-"],
        "net_statuses": {
            "ACTIVE": {"text": "#000000", "bg": "#00FF00"},
            "CLOSED": {"text": "#00FF00", "bg": "#000000"},
            "STANDBY": {"text": "#FFFF00", "bg": "#000000"}
        },
        "scripts": [
            {
                "title": "Preamble",
                "text": "Welcome to the Amateur Radio Emergency / Ragchew Net. This is [CALLSIGN], acting as Net Control Station for tonight.\n\nIs there any emergency or priority traffic on frequency? Please call now."
            }
        ],
        "api_sync": {
            "enabled": False,
            "url": "https://www.mcinares.org/chatgemilaude/doupdate/",
            "qrz": "",
            "param_name": "mcinaresstatus",
            "active_value": "ACTIVE",
            "inactive_value": "INACTIVE"
        }
    }
}

def load_station_memory():
    if not os.path.exists(STATION_MEMORY_FILE):
        return []
    try:
        with open(STATION_MEMORY_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
            if isinstance(data, list):
                out = []
                for item in data:
                    if isinstance(item, dict):
                        out.append(item)
                    elif isinstance(item, str):
                        out.append({"callsign": item, "name": ""})
                return out
            return []
    except Exception as e:
        print(f"Error loading station memory: {e}")
        return []

def save_station_memory(memory_list):
    try:
        unique_map = {}
        for item in memory_list:
            c = item.get("callsign", "").upper()
            if c:
                unique_map[c] = item.get("name", "")
        formatted = [{"callsign": k, "name": v} for k, v in unique_map.items()]
        with open(STATION_MEMORY_FILE, "w", encoding="utf-8") as f:
            json.dump(formatted, f, indent=2)
    except Exception as e:
        print(f"Error saving station memory: {e}")

def generate_stylesheet(font_size=10):
    return f"""
QMainWindow, QWidget, QDialog {{
    background-color: #C0C0C0;
    color: #000000;
    font-family: 'Tahoma', 'MS Sans Serif', sans-serif;
    font-size: {font_size}pt;
}}
QGroupBox {{
    border: 2px solid;
    border-color: #FFFFFF #808080 #808080 #FFFFFF;
    margin-top: 22px;
    font-weight: bold;
    color: #000080;
}}
QGroupBox::title {{
    subcontrol-origin: margin;
    subcontrol-position: top left;
    left: 8px;
    padding: 0 4px;
    background-color: #C0C0C0;
}}
QLineEdit, QComboBox, QTextEdit, QListWidget {{
    background-color: #FFFFFF;
    color: #000000;
    border: 2px solid;
    border-color: #808080 #FFFFFF #FFFFFF #808080;
    padding: 1px 2px;
    selection-background-color: #000080;
    selection-color: #FFFFFF;
}}
QCheckBox {{
    font-weight: bold;
    color: #000000;
}}
QPushButton {{
    background-color: #C0C0C0;
    color: #000000;
    border: 2px solid;
    border-color: #FFFFFF #808080 #808080 #FFFFFF;
    padding: 2px 4px;
    font-weight: bold;
}}
QPushButton:focus {{
    border: 2px dotted #000000;
}}
QPushButton:pressed {{
    border-color: #808080 #FFFFFF #FFFFFF #808080;
    padding: 3px 3px 1px 5px;
}}
QTableWidget {{
    background-color: #FFFFFF;
    color: #000000;
    gridline-color: #C0C0C0;
    border: 2px solid;
    border-color: #808080 #FFFFFF #FFFFFF #808080;
}}
QTableWidget::item:selected, QListWidget::item:selected {{
    background-color: #000080 !important;
    color: #FFFFFF !important;
}}
QHeaderView::section {{
    background-color: #C0C0C0;
    color: #000000;
    padding: 1px 2px;
    border: 2px solid;
    border-color: #FFFFFF #808080 #808080 #FFFFFF;
    font-weight: bold;
}}
QTabWidget::pane {{
    border: 2px solid;
    border-color: #FFFFFF #808080 #808080 #FFFFFF;
    background-color: #C0C0C0;
}}
QTabBar::tab {{
    background-color: #C0C0C0;
    color: #000000;
    padding: 2px 6px;
    border: 2px solid;
    border-color: #FFFFFF #808080 #C0C0C0 #FFFFFF;
    margin-right: 2px;
}}
QTabBar::tab:selected {{
    border-color: #FFFFFF #808080 #C0C0C0 #FFFFFF;
    font-weight: bold;
    background-color: #C0C0C0;
}}
"""

BTN_START_STYLE = """
QPushButton {
    background-color: #008000;
    color: #FFFFFF;
    border: 2px solid;
    border-color: #FFFFFF #004000 #004000 #FFFFFF;
    padding: 2px 4px;
    font-weight: bold;
}
QPushButton:pressed {
    border-color: #004000 #FFFFFF #FFFFFF #004000;
    padding: 3px 3px 1px 5px;
}
"""

BTN_GREEN_STYLE = BTN_START_STYLE

BTN_CLOSE_STYLE = """
QPushButton {
    background-color: #CC0000;
    color: #FFFFFF;
    border: 2px solid;
    border-color: #FFFFFF #660000 #660000 #FFFFFF;
    padding: 2px 4px;
    font-weight: bold;
}
QPushButton:pressed {
    border-color: #660000 #FFFFFF #FFFFFF #660000;
    padding: 3px 3px 1px 5px;
}
"""

BTN_ORANGE_STYLE = """
QPushButton {
    background-color: #FF8C00;
    color: #FFFFFF;
    border: 2px solid;
    border-color: #FFFFFF #8B4500 #8B4500 #FFFFFF;
    padding: 2px 4px;
    font-weight: bold;
}
QPushButton:pressed {
    border-color: #8B4500 #FFFFFF #FFFFFF #8B4500;
    padding: 3px 3px 1px 5px;
}
"""

BTN_CYAN_STYLE = """
QPushButton {
    background-color: #00FFFF;
    color: #000000;
    border: 2px solid;
    border-color: #FFFFFF #008080 #008080 #FFFFFF;
    padding: 2px 4px;
    font-weight: bold;
}
QPushButton:pressed {
    border-color: #008080 #FFFFFF #FFFFFF #008080;
}
"""

BTN_SAVE_STYLE = """
QPushButton {
    background-color: #CCCC00;
    color: #000000;
    border: 2px solid;
    border-color: #FFFFFF #666600 #666600 #FFFFFF;
    padding: 2px 4px;
    font-weight: bold;
}
QPushButton:pressed {
    border-color: #666600 #FFFFFF #FFFFFF #666600;
}
"""

BTN_RESTORE_STYLE = BTN_ORANGE_STYLE
BTN_DANGER_STYLE = BTN_CLOSE_STYLE

def load_profiles_from_disk():
    if not os.path.exists(PROFILES_FILE):
        save_profiles_to_disk(DEFAULT_PROFILES)
        return DEFAULT_PROFILES.copy()
    try:
        with open(PROFILES_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        print(f"Error loading profiles: {e}")
        return DEFAULT_PROFILES.copy()

def save_profiles_to_disk(profiles_data):
    try:
        with open(PROFILES_FILE, "w", encoding="utf-8") as f:
            json.dump(profiles_data, f, indent=2)
    except Exception as e:
        print(f"Error saving profiles: {e}")

def parse_color_map(val):
    if isinstance(val, dict):
        return val.get("text", "#FFFF00"), val.get("bg", "#000000")
    elif isinstance(val, str) and val.startswith("#"):
        return "#000000", val
    return "#FFFF00", "#000000"

def dispatch_api_status_update(api_cfg, status_text_or_flag, app_ref=None):
    if not api_cfg or not api_cfg.get("enabled", False):
        return

    def worker():
        url = api_cfg.get("url", "").strip()
        qrz = api_cfg.get("qrz", "").strip()
        param = api_cfg.get("param_name", "").strip()
        active_val = api_cfg.get("active_value", "ACTIVE").strip() or "ACTIVE"
        inactive_val = api_cfg.get("inactive_value", "INACTIVE").strip() or "INACTIVE"

        if not url or not qrz or not param:
            return

        if isinstance(status_text_or_flag, bool):
            post_val = active_val if status_text_or_flag else inactive_val
        else:
            txt = str(status_text_or_flag).strip()
            txt_upper = txt.upper()
            if txt_upper in ("ACTIVE", "LEVEL 1", "LEVEL 2", "LEVEL 3", "WARNING", "WATCH"):
                post_val = active_val
            elif txt_upper in ("INACTIVE", "CLOSED", "STANDBY"):
                post_val = inactive_val
            else:
                post_val = txt

        payload = {
            "qrz": qrz,
            param: post_val
        }

        try:
            encoded_data = urllib.parse.urlencode(payload).encode("utf-8")
            req = urllib.request.Request(url, data=encoded_data, method="POST")
            req.add_header("User-Agent", "NCM4-DesktopApp")

            with urllib.request.urlopen(req, timeout=6) as resp:
                status_code = resp.status
        except Exception as e:
            if app_ref:
                app_ref.add_activity_log("API ERROR", "HTTP FAIL", "-", f"Failed to sync {param}: {e}")

    threading.Thread(target=worker, daemon=True).start()


class CallsignCompletionModel(QAbstractItemModel):
    def __init__(self, memory_list, parent=None):
        super().__init__(parent)
        self.items = memory_list

    def set_items(self, memory_list):
        self.beginResetModel()
        self.items = memory_list
        self.endResetModel()

    def rowCount(self, parent=QModelIndex()):
        return len(self.items) if not parent.isValid() else 0

    def columnCount(self, parent=QModelIndex()):
        return 1

    def index(self, row, column, parent=QModelIndex()):
        if 0 <= row < len(self.items) and column == 0:
            return self.createIndex(row, column)
        return QModelIndex()

    def parent(self, index):
        return QModelIndex()

    def data(self, index, role=Qt.ItemDataRole.DisplayRole):
        if not index.isValid() or index.row() >= len(self.items):
            return None
        item = self.items[index.row()]
        call = item.get("callsign", "")
        name = item.get("name", "")

        if role == Qt.ItemDataRole.DisplayRole:
            return f"{call} — {name}" if name else call
        elif role == Qt.ItemDataRole.EditRole:
            return call
        return None


class UppercaseLineEdit(QLineEdit):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.textChanged.connect(self.handle_text_changed)

    def handle_text_changed(self, text):
        upper_text = text.upper()
        if text != upper_text:
            pos = self.cursorPosition()
            self.setText(upper_text)
            self.setCursorPosition(pos)


class EnterToggleCheckBox(QCheckBox):
    def keyPressEvent(self, event):
        if event.key() in (Qt.Key.Key_Return, Qt.Key.Key_Enter):
            self.animateClick()
        else:
            super().keyPressEvent(event)


class EnterClickButton(QPushButton):
    def keyPressEvent(self, event):
        if event.key() in (Qt.Key.Key_Return, Qt.Key.Key_Enter):
            self.animateClick()
        else:
            super().keyPressEvent(event)


class ScriptTextEdit(QTextEdit):
    def keyPressEvent(self, event):
        if event.key() in (Qt.Key.Key_Return, Qt.Key.Key_Enter):
            if event.modifiers() & Qt.KeyboardModifier.ShiftModifier:
                self.insertPlainText("\n")
                return
        super().keyPressEvent(event)


class NcsScribeLoginDialog(QDialog):
    def __init__(self, default_ncs="", station_types=None, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Net Control Station & Scribe Login")
        self.setMinimumWidth(380)

        self.station_types = station_types or ["Base", "Mobile", "NCS", "Alternate NCS"]

        layout = QVBoxLayout(self)

        # Role Selection Dropdown
        lbl_role = QLabel("Select Operating Role:")
        lbl_role.setFont(QFont("Tahoma", 9, QFont.Weight.Bold))
        layout.addWidget(lbl_role)

        self.combo_role = QComboBox()
        self.combo_role.addItems(["NCS (Net Control Station)", "ANCS (Alternate Net Control)", "Scribe (Assistant)"])
        layout.addWidget(self.combo_role)

        lbl_ncs = QLabel("Enter Operator Callsign:")
        lbl_ncs.setFont(QFont("Tahoma", 9, QFont.Weight.Bold))
        layout.addWidget(lbl_ncs)

        self.input_ncs = UppercaseLineEdit()
        self.input_ncs.setPlaceholderText("e.g. W9A")
        self.input_ncs.setText(default_ncs)
        layout.addWidget(self.input_ncs)

        self.chk_enable_scribe = EnterToggleCheckBox("Log in additional Scribe Callsign")
        self.chk_enable_scribe.toggled.connect(self.toggle_scribe_container)
        layout.addWidget(self.chk_enable_scribe)

        self.scribe_container = QWidget()
        scribe_layout = QGridLayout(self.scribe_container)
        scribe_layout.setContentsMargins(0, 4, 0, 4)

        lbl_scribe = QLabel("Enter Scribe / Assistant Callsign:")
        lbl_scribe.setFont(QFont("Tahoma", 9, QFont.Weight.Bold))
        scribe_layout.addWidget(lbl_scribe, 0, 0, 1, 2)

        self.input_scribe = UppercaseLineEdit()
        self.input_scribe.setPlaceholderText("e.g. K9B")
        scribe_layout.addWidget(self.input_scribe, 1, 0)

        self.combo_scribe_type = QComboBox()
        self.combo_scribe_type.addItems(self.station_types)
        
        matched_idx = 0
        for idx, st_type in enumerate(self.station_types):
            if "ALTERNATE" in st_type.upper() or "SCRIBE" in st_type.upper():
                matched_idx = idx
                break
        self.combo_scribe_type.setCurrentIndex(matched_idx)

        scribe_layout.addWidget(self.combo_scribe_type, 1, 1)

        layout.addWidget(self.scribe_container)
        self.scribe_container.setVisible(False)

        btn_layout = QHBoxLayout()
        btn_confirm = EnterClickButton("Log In")
        btn_confirm.setStyleSheet(BTN_GREEN_STYLE)
        btn_confirm.clicked.connect(self.handle_accept)

        btn_cancel = QPushButton("Cancel")
        btn_cancel.clicked.connect(self.reject)

        btn_layout.addStretch()
        btn_layout.addWidget(btn_confirm)
        btn_layout.addWidget(btn_cancel)
        layout.addLayout(btn_layout)

    def toggle_scribe_container(self, enabled):
        self.scribe_container.setVisible(enabled)
        self.adjustSize()

    def handle_accept(self):
        ncs_val = self.input_ncs.text().strip()
        scribe_val = self.input_scribe.text().strip()

        if not ncs_val:
            reply = QMessageBox.question(
                self, "Missing Callsign",
                "Operator Callsign is empty. Continue with 'UNASSIGNED'?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
            )
            if reply == QMessageBox.StandardButton.No:
                return

        if self.chk_enable_scribe.isChecked() and not scribe_val:
            reply = QMessageBox.question(
                self, "Missing Scribe Callsign",
                "Scribe option is enabled but field is blank. Continue without Scribe?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
            )
            if reply == QMessageBox.StandardButton.No:
                return

        self.accept()

    def get_logins(self):
        ncs = self.input_ncs.text().strip().upper() or "UNASSIGNED"
        scribe = ""
        scribe_type = self.combo_scribe_type.currentText()
        if self.chk_enable_scribe.isChecked():
            scribe = self.input_scribe.text().strip().upper()
        role_text = self.combo_role.currentText().split()[0]
        return ncs, scribe, scribe_type, role_text


class CombinedStationSearchEditDialog(QDialog):
    def __init__(self, stations, station_types, is_skywarn=False, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Search & Edit Station Roster")
        self.setMinimumSize(540, 380)

        self.stations = stations
        self.station_types = station_types
        self.is_skywarn = is_skywarn
        self.matching_indices = []
        self.current_match_pos = -1

        layout = QVBoxLayout(self)

        search_gb = QGroupBox("Search Roster Stations")
        search_layout = QGridLayout(search_gb)

        search_layout.addWidget(QLabel("Callsign / Substring:"), 0, 0)
        self.input_search_call = UppercaseLineEdit()
        search_layout.addWidget(self.input_search_call, 0, 1)

        search_layout.addWidget(QLabel("Station Type:"), 0, 2)
        self.combo_search_type = QComboBox()
        self.combo_search_type.addItem("-- Any --")
        self.combo_search_type.addItems(station_types)
        search_layout.addWidget(self.combo_search_type, 0, 3)

        search_layout.addWidget(QLabel("Status:"), 1, 0)
        self.combo_search_status = QComboBox()
        self.combo_search_status.addItems(["-- Any --", "Active", "Standby", "Checked Out"])
        search_layout.addWidget(self.combo_search_status, 1, 1)

        btn_find = EnterClickButton("Search")
        btn_find.clicked.connect(self.execute_search)
        search_layout.addWidget(btn_find, 1, 3)

        layout.addWidget(search_gb)

        nav_row = QHBoxLayout()
        self.lbl_match_count = QLabel("Matches Found: 0")
        self.lbl_match_count.setFont(QFont("Tahoma", 9, QFont.Weight.Bold))

        self.btn_prev_match = QPushButton("◄ Previous Match")
        self.btn_prev_match.clicked.connect(self.prev_match)

        self.btn_next_match = QPushButton("Next Match ►")
        self.btn_next_match.clicked.connect(self.next_match)

        nav_row.addWidget(self.lbl_match_count)
        nav_row.addStretch()
        nav_row.addWidget(self.btn_prev_match)
        nav_row.addWidget(self.btn_next_match)
        layout.addLayout(nav_row)

        edit_gb = QGroupBox("Edit Selected Station")
        edit_grid = QGridLayout(edit_gb)

        edit_grid.addWidget(QLabel("Callsign:"), 0, 0)
        self.edit_call = UppercaseLineEdit()
        edit_grid.addWidget(self.edit_call, 0, 1)

        edit_grid.addWidget(QLabel("Station Type:"), 0, 2)
        self.edit_type = QComboBox()
        self.edit_type.addItems(station_types)
        edit_grid.addWidget(self.edit_type, 0, 3)

        if not is_skywarn:
            self.chk_traffic = EnterToggleCheckBox("Traffic")
            self.chk_announcement = EnterToggleCheckBox("Announcement")
            chk_layout = QHBoxLayout()
            chk_layout.addWidget(self.chk_traffic)
            chk_layout.addWidget(self.chk_announcement)
            edit_grid.addLayout(chk_layout, 1, 0, 1, 2)

            edit_grid.addWidget(QLabel("Info / Grid:"), 1, 2)
            self.edit_info = QLineEdit()
            edit_grid.addWidget(self.edit_info, 1, 3)
        else:
            self.chk_traffic = None
            self.chk_announcement = None

            edit_grid.addWidget(QLabel("Location:"), 1, 0)
            self.edit_info = QLineEdit()
            edit_grid.addWidget(self.edit_info, 1, 1, 1, 3)

        edit_grid.addWidget(QLabel("Status:"), 2, 0)
        self.edit_status = QComboBox()
        self.edit_status.addItems(["Active", "Standby", "Checked Out"])
        edit_grid.addWidget(self.edit_status, 2, 1)

        layout.addWidget(edit_gb, 1)

        btn_layout = QHBoxLayout()
        btn_save = EnterClickButton("Save Edits to Station")
        btn_save.setStyleSheet(BTN_GREEN_STYLE)
        btn_save.clicked.connect(self.save_current_station)

        btn_close = QPushButton("Close")
        btn_close.clicked.connect(self.accept)

        btn_layout.addStretch()
        btn_layout.addWidget(btn_save)
        btn_layout.addWidget(btn_close)
        layout.addLayout(btn_layout)

    def execute_search(self):
        target_call = self.input_search_call.text().strip().upper()
        target_type = self.combo_search_type.currentText()
        target_status = self.combo_search_status.currentText()

        self.matching_indices = []

        for idx, st in enumerate(self.stations):
            c = st.get("callsign", "").upper()
            t = st.get("type", "")
            s = st.get("status", "")

            if target_call and target_call not in c:
                continue
            if target_type != "-- Any --" and target_type != t:
                continue
            if target_status != "-- Any --" and target_status != s:
                continue

            self.matching_indices.append(idx)

        self.lbl_match_count.setText(f"Matches Found: {len(self.matching_indices)}")
        if self.matching_indices:
            self.current_match_pos = 0
            self.load_match_data()
        else:
            self.current_match_pos = -1
            self.clear_edit_fields()

    def load_match_data(self):
        if 0 <= self.current_match_pos < len(self.matching_indices):
            idx = self.matching_indices[self.current_match_pos]
            st = self.stations[idx]

            self.edit_call.setText(st.get("callsign", ""))
            
            type_text = st.get("type", "")
            if self.edit_type.findText(type_text) >= 0:
                self.edit_type.setCurrentText(type_text)
            
            self.edit_status.setCurrentText(st.get("status", "Active"))

            if self.is_skywarn:
                self.edit_info.setText(st.get("detail", ""))
            else:
                self.edit_info.setText(st.get("info", ""))
                detail_str = st.get("detail", "")
                if self.chk_traffic:
                    self.chk_traffic.setChecked("Traffic" in detail_str)
                if self.chk_announcement:
                    self.chk_announcement.setChecked("Announcement" in detail_str)

            self.lbl_match_count.setText(f"Match {self.current_match_pos + 1} of {len(self.matching_indices)}")

    def save_current_station(self):
        if not self.matching_indices or not (0 <= self.current_match_pos < len(self.matching_indices)):
            QMessageBox.information(self, "No Station Selected", "There is no station currently loaded to edit.")
            return

        idx = self.matching_indices[self.current_match_pos]
        st = self.stations[idx]

        call = self.edit_call.text().strip().upper()
        if not call:
            QMessageBox.warning(self, "Validation Error", "Callsign cannot be empty.")
            return

        st["callsign"] = call
        st["type"] = self.edit_type.currentText()
        st["status"] = self.edit_status.currentText()

        info_val = self.edit_info.text().strip() or "-"
        if self.is_skywarn:
            st["detail"] = info_val
        else:
            st["info"] = info_val
            items = []
            if self.chk_traffic and self.chk_traffic.isChecked():
                items.append("Traffic")
            if self.chk_announcement and self.chk_announcement.isChecked():
                items.append("Announcement")
            st["detail"] = ", ".join(items) if items else "-"

        QMessageBox.information(self, "Station Saved", f"Station '{call}' updated successfully.")

    def prev_match(self):
        if self.matching_indices and self.current_match_pos > 0:
            self.current_match_pos -= 1
            self.load_match_data()

    def next_match(self):
        if self.matching_indices and self.current_match_pos < len(self.matching_indices) - 1:
            self.current_match_pos += 1
            self.load_match_data()

    def clear_edit_fields(self):
        self.edit_call.clear()
        self.edit_info.clear()
        if self.chk_traffic:
            self.chk_traffic.setChecked(False)
        if self.chk_announcement:
            self.chk_announcement.setChecked(False)


class CombinedTrafficSearchEditDialog(QDialog):
    def __init__(self, activities, traffic_types, is_skywarn=False, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Search & Edit Net Traffic Log")
        self.setMinimumSize(540, 420)

        self.activities = activities
        self.traffic_types = traffic_types
        self.is_skywarn = is_skywarn
        self.matching_indices = []
        self.current_match_pos = -1

        layout = QVBoxLayout(self)

        search_gb = QGroupBox("Search Log Entries")
        search_layout = QGridLayout(search_gb)

        search_layout.addWidget(QLabel("Callsign / Substring:"), 0, 0)
        self.input_search_call = UppercaseLineEdit()
        search_layout.addWidget(self.input_search_call, 0, 1)

        search_layout.addWidget(QLabel("Type / Tag:"), 0, 2)
        self.combo_search_type = QComboBox()
        self.combo_search_type.addItem("-- Any --")
        self.combo_search_type.addItems(traffic_types)
        search_layout.addWidget(self.combo_search_type, 0, 3)

        search_layout.addWidget(QLabel("Timestamp:"), 1, 0)
        self.input_search_time = QLineEdit()
        self.input_search_time.setPlaceholderText("HH:MM")
        search_layout.addWidget(self.input_search_time, 1, 1)

        search_layout.addWidget(QLabel("Time Match Rule:"), 1, 2)
        self.combo_time_op = QComboBox()
        self.combo_time_op.addItems(["At", "At or Before", "At or After"])
        search_layout.addWidget(self.combo_time_op, 1, 3)

        btn_find = QPushButton("Search")
        btn_find.clicked.connect(self.execute_search)
        search_layout.addWidget(btn_find, 2, 3)

        layout.addWidget(search_gb)

        nav_row = QHBoxLayout()
        self.lbl_match_count = QLabel("Matches Found: 0")
        self.lbl_match_count.setFont(QFont("Tahoma", 9, QFont.Weight.Bold))

        self.btn_prev_match = QPushButton("◄ Previous Match")
        self.btn_prev_match.clicked.connect(self.prev_match)

        self.btn_next_match = QPushButton("Next Match ►")
        self.btn_next_match.clicked.connect(self.next_match)

        nav_row.addWidget(self.lbl_match_count)
        nav_row.addStretch()
        nav_row.addWidget(self.btn_prev_match)
        nav_row.addWidget(self.btn_next_match)
        layout.addLayout(nav_row)

        edit_gb = QGroupBox("Edit Selected Entry")
        edit_grid = QGridLayout(edit_gb)

        edit_grid.addWidget(QLabel("Station Callsign:"), 0, 0)
        self.edit_call = UppercaseLineEdit()
        edit_grid.addWidget(self.edit_call, 0, 1)

        edit_grid.addWidget(QLabel("Traffic Type:"), 0, 2)
        self.edit_type = QComboBox()
        self.edit_type.addItems(traffic_types)
        edit_grid.addWidget(self.edit_type, 0, 3)

        edit_grid.addWidget(QLabel("Time (HH:MM):"), 1, 0)
        self.edit_time = QLineEdit()
        edit_grid.addWidget(self.edit_time, 1, 1)

        if is_skywarn:
            edit_grid.addWidget(QLabel("Location:"), 1, 2)
            self.edit_location = QLineEdit()
            edit_grid.addWidget(self.edit_location, 1, 3)

        edit_grid.addWidget(QLabel("Details / Note:"), 2, 0)
        self.edit_note = QTextEdit()
        edit_grid.addWidget(self.edit_note, 2, 1, 1, 3)

        layout.addWidget(edit_gb, 1)

        btn_layout = QHBoxLayout()
        btn_save = EnterClickButton("Save Edits to Entry")
        btn_save.setStyleSheet(BTN_GREEN_STYLE)
        btn_save.clicked.connect(self.save_current_entry)

        btn_close = QPushButton("Close")
        btn_close.clicked.connect(self.accept)

        btn_layout.addStretch()
        btn_layout.addWidget(btn_save)
        btn_layout.addWidget(btn_close)
        layout.addLayout(btn_layout)

    def execute_search(self):
        target_call = self.input_search_call.text().strip().upper()
        target_type = self.combo_search_type.currentText()
        target_time = self.input_search_time.text().strip()
        time_op = self.combo_time_op.currentText()

        self.matching_indices = []

        for idx, act in enumerate(self.activities):
            c = act.get("callsign", act.get("source", "")).upper()
            t = act.get("type", act.get("condition", ""))
            ts = act.get("timestamp", "")

            if target_call and target_call not in c:
                continue

            if target_type != "-- Any --" and target_type != t:
                continue

            if target_time:
                if time_op == "At" and ts != target_time:
                    continue
                elif time_op == "At or Before" and ts > target_time:
                    continue
                elif time_op == "At or After" and ts < target_time:
                    continue

            self.matching_indices.append(idx)

        self.lbl_match_count.setText(f"Matches Found: {len(self.matching_indices)}")
        if self.matching_indices:
            self.current_match_pos = 0
            self.load_match_data()
        else:
            self.current_match_pos = -1
            self.clear_edit_fields()

    def load_match_data(self):
        if 0 <= self.current_match_pos < len(self.matching_indices):
            idx = self.matching_indices[self.current_match_pos]
            act = self.activities[idx]

            self.edit_call.setText(act.get("callsign", act.get("source", "")))
            self.edit_type.setCurrentText(act.get("type", act.get("condition", "")))
            self.edit_time.setText(act.get("timestamp", ""))
            if self.is_skywarn:
                self.edit_location.setText(act.get("location", "-"))
            self.edit_note.setText(act.get("note", act.get("narrative", "")))

            self.lbl_match_count.setText(f"Match {self.current_match_pos + 1} of {len(self.matching_indices)}")

    def save_current_entry(self):
        if 0 <= self.current_match_pos < len(self.matching_indices):
            idx = self.matching_indices[self.current_match_pos]
            act = self.activities[idx]

            if self.is_skywarn:
                act["source"] = self.edit_call.text().strip().upper()
                act["condition"] = self.edit_type.currentText()
                act["timestamp"] = self.edit_time.text().strip()
                act["location"] = self.edit_location.text().strip() or "-"
                act["narrative"] = self.edit_note.toPlainText().strip()
            else:
                act["callsign"] = self.edit_call.text().strip().upper()
                act["type"] = self.edit_type.currentText()
                act["timestamp"] = self.edit_time.text().strip()
                act["note"] = self.edit_note.toPlainText().strip()

            QMessageBox.information(self, "Entry Saved", "Log entry updated successfully.")

    def prev_match(self):
        if self.matching_indices and self.current_match_pos > 0:
            self.current_match_pos -= 1
            self.load_match_data()

    def next_match(self):
        if self.matching_indices and self.current_match_pos < len(self.matching_indices) - 1:
            self.current_match_pos += 1
            self.load_match_data()

    def clear_edit_fields(self):
        self.edit_call.clear()
        self.edit_time.clear()
        if self.is_skywarn:
            self.edit_location.clear()
        self.edit_note.clear()


class ProfileAuthDialog(QDialog):
    def __init__(self, profile_name, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Profile Access Authentication")
        self.setFixedSize(380, 180)

        layout = QVBoxLayout(self)
        lbl = QLabel(f"Profile '{profile_name}' is passcode protected.\nEnter authorized Callsign and Passcode:")
        lbl.setFont(QFont("Tahoma", 9, QFont.Weight.Bold))
        lbl.setWordWrap(True)
        layout.addWidget(lbl)

        grid = QGridLayout()
        grid.addWidget(QLabel("Callsign:"), 0, 0)
        self.input_call = UppercaseLineEdit()
        grid.addWidget(self.input_call, 0, 1)

        grid.addWidget(QLabel("Passcode / PIN:"), 1, 0)
        self.input_pass = QLineEdit()
        self.input_pass.setEchoMode(QLineEdit.EchoMode.Password)
        self.input_pass.returnPressed.connect(self.accept)
        grid.addWidget(self.input_pass, 1, 1)

        layout.addLayout(grid)

        btn_layout = QHBoxLayout()
        btn_ok = EnterClickButton("Unlock Profile")
        btn_ok.setStyleSheet(BTN_GREEN_STYLE)
        btn_ok.clicked.connect(self.accept)

        btn_cancel = QPushButton("Cancel")
        btn_cancel.clicked.connect(self.reject)

        btn_layout.addStretch()
        btn_layout.addWidget(btn_ok)
        btn_layout.addWidget(btn_cancel)
        layout.addLayout(btn_layout)

    def get_credentials(self):
        return self.input_call.text().strip().upper(), self.input_pass.text().strip()


class UserManagementDialog(QDialog):
    def __init__(self, users_data, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Access Control - User Management")
        self.setMinimumSize(460, 320)
        self.users = json.loads(json.dumps(users_data))

        layout = QVBoxLayout(self)

        lbl_info = QLabel("Authorized Net Control Operators & Custom Passcodes:")
        lbl_info.setFont(QFont("Tahoma", 9, QFont.Weight.Bold))
        layout.addWidget(lbl_info)

        self.table_users = QTableWidget(0, 2)
        self.table_users.setHorizontalHeaderLabels(["Operator Callsign", "Custom Passcode / PIN"])
        self.table_users.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        self.table_users.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)

        for user in self.users:
            row = self.table_users.rowCount()
            self.table_users.insertRow(row)

            call_item = QTableWidgetItem(user.get("callsign", ""))
            pass_item = QTableWidgetItem(user.get("passcode", ""))

            self.table_users.setItem(row, 0, call_item)
            self.table_users.setItem(row, 1, pass_item)

        layout.addWidget(self.table_users, 1)

        tbl_btns = QHBoxLayout()
        btn_add = QPushButton("Add User")
        btn_add.clicked.connect(self.add_user_row)
        btn_del = QPushButton("Remove Selected User")
        btn_del.clicked.connect(self.delete_user_row)

        tbl_btns.addWidget(btn_add)
        tbl_btns.addWidget(btn_del)
        tbl_btns.addStretch()
        layout.addLayout(tbl_btns)

        bottom_btns = QHBoxLayout()
        btn_save = EnterClickButton("Save Users")
        btn_save.setStyleSheet(BTN_GREEN_STYLE)
        btn_save.clicked.connect(self.accept)

        btn_cancel = QPushButton("Cancel")
        btn_cancel.clicked.connect(self.reject)

        bottom_btns.addStretch()
        bottom_btns.addWidget(btn_save)
        bottom_btns.addWidget(btn_cancel)
        layout.addLayout(bottom_btns)

    def add_user_row(self):
        row = self.table_users.rowCount()
        self.table_users.insertRow(row)
        self.table_users.setItem(row, 0, QTableWidgetItem("W9A"))
        self.table_users.setItem(row, 1, QTableWidgetItem("1234"))

    def delete_user_row(self):
        curr = self.table_users.currentRow()
        if curr >= 0:
            self.table_users.removeRow(curr)

    def get_users_data(self):
        out = []
        for r in range(self.table_users.rowCount()):
            c_item = self.table_users.item(r, 0)
            p_item = self.table_users.item(r, 1)

            call = c_item.text().strip().upper() if c_item else ""
            passc = p_item.text().strip() if p_item else ""

            if call:
                out.append({"callsign": call, "passcode": passc})
        return out


class RepeaterInfoDialog(QDialog):
    def __init__(self, repeater_data, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Repeater Information Configuration")
        self.setMinimumWidth(400)

        layout = QVBoxLayout(self)

        # Primary Repeater Group
        primary_gb = QGroupBox("Primary Repeater")
        p_grid = QGridLayout(primary_gb)

        p_grid.addWidget(QLabel("Frequency:"), 0, 0)
        self.input_rep_freq = QLineEdit()
        self.input_rep_freq.setText(repeater_data.get("freq", "146.700"))
        p_grid.addWidget(self.input_rep_freq, 0, 1)

        p_grid.addWidget(QLabel("Offset:"), 1, 0)
        self.combo_rep_offset = QComboBox()
        self.combo_rep_offset.addItems(["-", "+", "Simplex", "None"])
        self.combo_rep_offset.setCurrentText(repeater_data.get("offset", "-"))
        p_grid.addWidget(self.combo_rep_offset, 1, 1)

        p_grid.addWidget(QLabel("PL Tone:"), 2, 0)
        self.input_rep_tone = QLineEdit()
        self.input_rep_tone.setText(repeater_data.get("tone", "88.5 PL"))
        p_grid.addWidget(self.input_rep_tone, 2, 1)

        layout.addWidget(primary_gb)

        # Backup Repeater Section
        has_backup = repeater_data.get("has_backup", False)
        self.chk_enable_backup = EnterToggleCheckBox("Enable Backup Repeater")
        self.chk_enable_backup.setChecked(has_backup)
        self.chk_enable_backup.toggled.connect(self.toggle_backup_fields)
        layout.addWidget(self.chk_enable_backup)

        self.backup_gb = QGroupBox("Backup Repeater")
        b_grid = QGridLayout(self.backup_gb)

        b_grid.addWidget(QLabel("Frequency:"), 0, 0)
        self.input_backup_freq = QLineEdit()
        self.input_backup_freq.setText(repeater_data.get("backup_freq", ""))
        b_grid.addWidget(self.input_backup_freq, 0, 1)

        b_grid.addWidget(QLabel("Offset:"), 1, 0)
        self.combo_backup_offset = QComboBox()
        self.combo_backup_offset.addItems(["-", "+", "Simplex", "None"])
        self.combo_backup_offset.setCurrentText(repeater_data.get("backup_offset", "-"))
        b_grid.addWidget(self.combo_backup_offset, 1, 1)

        b_grid.addWidget(QLabel("PL Tone:"), 2, 0)
        self.input_backup_tone = QLineEdit()
        self.input_backup_tone.setText(repeater_data.get("backup_tone", ""))
        b_grid.addWidget(self.input_backup_tone, 2, 1)

        layout.addWidget(self.backup_gb)
        self.backup_gb.setVisible(has_backup)

        btn_layout = QHBoxLayout()
        btn_save = EnterClickButton("Save Repeater Info")
        btn_save.setStyleSheet(BTN_GREEN_STYLE)
        btn_save.clicked.connect(self.accept)

        btn_cancel = QPushButton("Cancel")
        btn_cancel.clicked.connect(self.reject)

        btn_layout.addStretch()
        btn_layout.addWidget(btn_save)
        btn_layout.addWidget(btn_cancel)
        layout.addLayout(btn_layout)

    def toggle_backup_fields(self, enabled):
        self.backup_gb.setVisible(enabled)
        self.adjustSize()

    def get_repeater_data(self):
        return {
            "freq": self.input_rep_freq.text().strip() or "146.700",
            "offset": self.combo_rep_offset.currentText(),
            "tone": self.input_rep_tone.text().strip() or "88.5 PL",
            "has_backup": self.chk_enable_backup.isChecked(),
            "backup_freq": self.input_backup_freq.text().strip(),
            "backup_offset": self.combo_backup_offset.currentText(),
            "backup_tone": self.input_backup_tone.text().strip()
        }


class ApiConfigDialog(QDialog):
    def __init__(self, api_data, web_opts_data, mode_text="Regular", parent=None):
        super().__init__(parent)
        self.setWindowTitle("Configure Web Status API Sync")
        self.setMinimumSize(520, 360)

        self.api_data = json.loads(json.dumps(api_data))
        self.web_opts_data = list(web_opts_data)

        layout = QVBoxLayout(self)

        grid = QGridLayout()
        grid.addWidget(QLabel("Endpoint URL:"), 0, 0)
        self.input_url = QLineEdit()
        self.input_url.setText(self.api_data.get("url", ""))
        grid.addWidget(self.input_url, 0, 1)

        grid.addWidget(QLabel("Status Field Name:"), 1, 0)
        self.input_param = QLineEdit()
        self.input_param.setText(self.api_data.get("param_name", "skywarnstatus" if mode_text == "SKYWARN" else "mcinaresstatus"))
        grid.addWidget(self.input_param, 1, 1)

        grid.addWidget(QLabel("Default NCS Callsign:"), 2, 0)
        self.input_qrz = UppercaseLineEdit()
        self.input_qrz.setPlaceholderText("e.g. W9A")  # <-- Ghost text placeholder
        self.input_qrz.setText(self.api_data.get("qrz", "")) # <-- Empty fallback
        grid.addWidget(self.input_qrz, 2, 1)

        grid.addWidget(QLabel("Active Web Text:"), 3, 0)
        self.input_act = QLineEdit()
        self.input_act.setText(self.api_data.get("active_value", "ACTIVE"))
        grid.addWidget(self.input_act, 3, 1)

        grid.addWidget(QLabel("Inactive Web Text:"), 4, 0)
        self.input_inact = QLineEdit()
        self.input_inact.setText(self.api_data.get("inactive_value", "INACTIVE"))
        grid.addWidget(self.input_inact, 4, 1)

        layout.addLayout(grid)

        lbl_opts = QLabel("Custom Web Status Dropdown Options (One per line):")
        lbl_opts.setFont(QFont("Tahoma", 9, QFont.Weight.Bold))
        layout.addWidget(lbl_opts)

        self.text_opts = QTextEdit()
        self.text_opts.setText("\n".join(self.web_opts_data))
        layout.addWidget(self.text_opts, 1)

        btn_layout = QHBoxLayout()
        btn_save = EnterClickButton("Save Configuration")
        btn_save.setStyleSheet(BTN_GREEN_STYLE)
        btn_save.clicked.connect(self.accept)

        btn_cancel = QPushButton("Cancel")
        btn_cancel.clicked.connect(self.reject)

        btn_layout.addStretch()
        btn_layout.addWidget(btn_save)
        btn_layout.addWidget(btn_cancel)
        layout.addLayout(btn_layout)

    def get_config_data(self):
        updated_api = {
            "enabled": self.api_data.get("enabled", True),
            "url": self.input_url.text().strip(),
            "qrz": self.input_qrz.text().strip().upper(),
            "param_name": self.input_param.text().strip(),
            "active_value": self.input_act.text().strip(),
            "inactive_value": self.input_inact.text().strip()
        }
        updated_opts = [o.strip() for o in self.text_opts.toPlainText().split("\n") if o.strip()]
        return updated_api, updated_opts


class HotkeyHelpDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Application Hotkey & Shortcut Reference")
        self.setMinimumSize(540, 420)

        layout = QVBoxLayout(self)

        tabs = QTabWidget()

        sections = [
            ("Global & System", [
                ("F1", "Display Hotkey Reference"),
                ("F2", "Select Net Profile"),
                ("F3", "Open Profile Creator & Editor"),
                ("F4", "Start / Close Net Session"),
                ("F5", "Transfer Net Control Station"),
                ("Ctrl+S", "Save Session State"),
                ("Ctrl+O", "Open Saved Session File"),
                ("Ctrl+R", "Reset Net Session"),
                ("Ctrl+ / Ctrl+=", "Zoom In UI Font Size"),
                ("Ctrl+- / Ctrl+_", "Zoom Out UI Font Size"),
                ("Alt+F4", "Exit Application")
            ]),
            ("Station Management", [
                ("Alt+1", "Focus Check-In Callsign Field"),
                ("F8", "Start Net Roll Call"),
                ("Ctrl+F", "Search & Edit Roster Stations"),
                ("Ctrl+T", "Toggle Selected Station Standby/Active"),
                ("Ctrl+U", "Check Out Selected Station"),
                ("Ctrl+Shift+Del", "Delete Selected Station from Roster")
            ]),
            ("Traffic & Activity Logging", [
                ("Alt+2", "Focus Traffic Callsign Field"),
                ("Ctrl+Enter", "Log Traffic / Activity Entry"),
                ("Ctrl+Shift+T", "Open ARRL NTS Radiogram Form"),
                ("Ctrl+N", "Insert Current Time (SKYWARN Mode)"),
                ("F12 / Esc", "Clear Traffic Form Fields"),
                ("Ctrl+Shift+F", "Search & Edit Traffic Log"),
                ("Ctrl+Delete", "Delete Selected Traffic Entry")
            ]),
            ("Scripts, Web API & Status", [
                ("Ctrl+]", "Next Script Tab"),
                ("Ctrl+[", "Previous Script Tab"),
                ("F9", "Focus Web Status Dropdown"),
                ("F10", "Send/Push Web Status to API"),
                ("Ctrl+Alt+S", "Apply Selected Net Status Badge")
            ]),
            ("Roll Call Window", [
                ("Ctrl+Space", "Check In Current Station"),
                ("Ctrl+Down", "Skip Current Station"),
                ("Ctrl+Up", "Go to Previous Station")
            ])
        ]

        for title, hotkeys in sections:
            tab = QWidget()
            t_layout = QVBoxLayout(tab)
            table = QTableWidget(len(hotkeys), 2)
            table.setHorizontalHeaderLabels(["Hotkey", "Action / Function"])
            table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Interactive)
            table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
            table.setColumnWidth(0, 140)

            for row, (hk, desc) in enumerate(hotkeys):
                table.setItem(row, 0, QTableWidgetItem(hk))
                table.setItem(row, 1, QTableWidgetItem(desc))

            t_layout.addWidget(table)
            tabs.addTab(tab, title)

        layout.addWidget(tabs)

        btn_close = QPushButton("Close")
        btn_close.clicked.connect(self.accept)
        btn_layout = QHBoxLayout()
        btn_layout.addStretch()
        btn_layout.addWidget(btn_close)
        layout.addLayout(btn_layout)


class RedDangerDialog(QDialog):
    def __init__(self, title, message, require_type_delete=False, parent=None):
        super().__init__(parent)
        self.setWindowTitle(title)
        self.setFixedSize(400, 180)
        self.require_type_delete = require_type_delete

        self.setStyleSheet("""
            QDialog {
                background-color: #B00000;
                color: #FFFFFF;
            }
            QLabel {
                color: #FFFFFF;
                font-weight: bold;
            }
            QLineEdit {
                background-color: #FFFFFF;
                color: #000000;
                font-weight: bold;
            }
            QPushButton {
                background-color: #FFFFFF;
                color: #B00000;
                font-weight: bold;
                border: 2px solid #800000;
                padding: 4px;
            }
            QPushButton:pressed {
                background-color: #D0D0D0;
            }
        """)

        layout = QVBoxLayout(self)
        lbl = QLabel(message)
        lbl.setWordWrap(True)
        layout.addWidget(lbl)

        if require_type_delete:
            lbl_type = QLabel("Type 'DELETE' to confirm:")
            self.input_confirm = UppercaseLineEdit()
            layout.addWidget(lbl_type)
            layout.addWidget(self.input_confirm)

        btn_layout = QHBoxLayout()
        btn_confirm = QPushButton("Confirm")
        btn_confirm.clicked.connect(self.handle_confirm)

        btn_cancel = QPushButton("Cancel")
        btn_cancel.clicked.connect(self.reject)

        btn_layout.addStretch()
        btn_layout.addWidget(btn_confirm)
        btn_layout.addWidget(btn_cancel)
        layout.addLayout(btn_layout)

    def handle_confirm(self):
        if self.require_type_delete:
            if self.input_confirm.text().strip() == "DELETE":
                self.accept()
            else:
                QMessageBox.warning(self, "Validation Failed", "You must type 'DELETE' exactly to confirm.")
        else:
            self.accept()


class ProfileBadgeLabel(QLabel):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setFrameShape(QFrame.Shape.Box)
        self.setFrameShadow(QFrame.Shadow.Sunken)
        self.setLineWidth(2)
        self.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.setFixedSize(260, 38)
        self.setWordWrap(True)

        mono_font = QFont("Consolas", 8, QFont.Weight.Bold)
        mono_font.setStyleHint(QFont.StyleHint.Monospace)
        self.setFont(mono_font)

        self.setStyleSheet("""
            QLabel {
                background-color: #000000;
                color: #00FFFF;
                border: 2px solid #808080;
                padding: 1px 2px;
                font-family: 'Consolas', 'Courier New', monospace;
            }
        """)
        self.set_profile("NONE")

    def set_profile(self, profile_name):
        disp = profile_name.upper()
        if len(disp) > 42:
            disp = disp[:39] + "..."
        self.setText(disp)


class WebApiStatusBadgeLabel(QLabel):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setFrameShape(QFrame.Shape.Box)
        self.setFrameShadow(QFrame.Shadow.Sunken)
        self.setLineWidth(2)
        self.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.setFixedSize(280, 38)
        self.setWordWrap(True)

        mono_font = QFont("Consolas", 8, QFont.Weight.Bold)
        mono_font.setStyleHint(QFont.StyleHint.Monospace)
        self.setFont(mono_font)

        self.setStyleSheet("""
            QLabel {
                background-color: #000000;
                color: #FFFFFF;
                border: 2px solid #808080;
                padding: 1px 2px;
                font-family: 'Consolas', 'Courier New', monospace;
            }
        """)
        self.set_status("WEB STATUS: STANDBY")

    def set_status(self, text_val):
        clean_text = re.sub(r'<[^>]+>', '', str(text_val)).strip().upper()
        if not clean_text or clean_text == "STATUS" or "404" in clean_text:
            clean_text = "ERR/OFFLINE"
        self.setText(f"WEB STATUS:\n{clean_text}")


class RepeaterBadgeLabel(QLabel):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setFrameShape(QFrame.Shape.Box)
        self.setFrameShadow(QFrame.Shadow.Sunken)
        self.setLineWidth(2)
        self.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.setFixedSize(125, 38)
        self.setCursor(Qt.CursorShape.PointingHandCursor)

        mono_font = QFont("Consolas", 8, QFont.Weight.Bold)
        mono_font.setStyleHint(QFont.StyleHint.Monospace)
        self.setFont(mono_font)

        self.repeater_data = {}
        self.is_using_backup = False
        self.set_repeater_info({})

    def set_repeater_info(self, repeater_data, is_using_backup=False):
        self.repeater_data = repeater_data or {}
        self.is_using_backup = is_using_backup
        self.update_display()

    def update_display(self):
        if self.is_using_backup and self.repeater_data.get("has_backup"):
            f_str = self.repeater_data.get("backup_freq", "").strip() or "146.700"
            o_str = self.repeater_data.get("backup_offset", "-").strip()
            t_str = self.repeater_data.get("backup_tone", "").strip() or "88.5 PL"

            self.setStyleSheet("""
                QLabel {
                    background-color: #FFFF00;
                    color: #000000;
                    border: 2px solid #808080;
                    padding: 1px 2px;
                    font-family: 'Consolas', 'Courier New', monospace;
                }
            """)
        else:
            f_str = self.repeater_data.get("freq", "").strip() or "146.700"
            o_str = self.repeater_data.get("offset", "-").strip()
            t_str = self.repeater_data.get("tone", "").strip() or "88.5 PL"

            self.setStyleSheet("""
                QLabel {
                    background-color: #000000;
                    color: #00FF00;
                    border: 2px solid #808080;
                    padding: 1px 2px;
                    font-family: 'Consolas', 'Courier New', monospace;
                }
            """)

        line1 = f"{f_str} {o_str}".strip()
        self.setText(f"{line1}\n{t_str}")

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            main_window = self.window()
            if hasattr(main_window, "handle_repeater_badge_click"):
                main_window.handle_repeater_badge_click()
        super().mousePressEvent(event)


class NcsBadgeLabel(QLabel):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setFrameShape(QFrame.Shape.Box)
        self.setFrameShadow(QFrame.Shadow.Sunken)
        self.setLineWidth(2)
        self.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.setFixedSize(130, 38)

        mono_font = QFont("Consolas", 8, QFont.Weight.Bold)
        mono_font.setStyleHint(QFont.StyleHint.Monospace)
        self.setFont(mono_font)

        self.setStyleSheet("""
            QLabel {
                background-color: #000000;
                color: #FFFFFF;
                border: 2px solid #808080;
                padding: 1px 2px;
                font-family: 'Consolas', 'Courier New', monospace;
            }
        """)
        self.set_ncs("UNASSIGNED")

    def set_ncs(self, callsign, scribe_call=""):
        disp_ncs = callsign.upper() if callsign else "UNASSIGNED"
        if scribe_call:
            self.setText(f"NCS: {disp_ncs}\nSCRIBE: {scribe_call.upper()}")
        else:
            self.setText(f"NCS: {disp_ncs}")


class StatusBadgeLabel(QLabel):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setFrameShape(QFrame.Shape.Box)
        self.setFrameShadow(QFrame.Shadow.Sunken)
        self.setLineWidth(2)
        self.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.setFixedSize(110, 38)

        mono_font = QFont("Consolas", 8, QFont.Weight.Bold)
        mono_font.setStyleHint(QFont.StyleHint.Monospace)
        self.setFont(mono_font)
        self.set_state("INACTIVE", "#FFFF00", "#000000")

    def set_state(self, state_name, text_hex="#FFFF00", bg_hex="#000000"):
        state_name = state_name.upper()
        self.setText(f"STATUS:\n{state_name}")

        fg = text_hex if text_hex else "#FFFF00"
        bg = bg_hex if bg_hex else "#000000"

        self.setStyleSheet(f"""
            QLabel {{
                background-color: {bg};
                color: {fg};
                border: 2px solid #808080;
                padding: 1px 2px;
                font-family: 'Consolas', 'Courier New', monospace;
            }}
        """)


class DynamicTimerLabel(QLabel):
    INITIAL_TIME_SECONDS = 600

    def __init__(self, parent=None):
        super().__init__(parent)
        self.time_remaining = self.INITIAL_TIME_SECONDS
        self.flash_state = False

        self.setFrameShape(QFrame.Shape.Box)
        self.setFrameShadow(QFrame.Shadow.Sunken)
        self.setLineWidth(2)
        self.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.setCursor(Qt.CursorShape.PointingHandCursor)
        self.setFixedSize(110, 38)

        mono_font = QFont("Consolas", 8, QFont.Weight.Bold)
        mono_font.setStyleHint(QFont.StyleHint.Monospace)
        self.setFont(mono_font)

        self.countdown_timer = QTimer(self)
        self.countdown_timer.setInterval(1000)
        self.countdown_timer.timeout.connect(self.update_countdown)

        self.flash_timer = QTimer(self)
        self.flash_timer.setInterval(1000)
        self.flash_timer.timeout.connect(self.toggle_flash)

        self.apply_normal_style()
        self.update_display_text()

    def update_countdown(self):
        if self.time_remaining > 0:
            self.time_remaining -= 1
            self.update_display_text()

            if self.time_remaining == 0:
                self.start_flashing()

    def update_display_text(self):
        mins = self.time_remaining // 60
        secs = self.time_remaining % 60
        self.setText(f"ID TIMER\n{mins:02d}:{secs:02d}")

    def start_flashing(self):
        self.setText("ID NOW")
        if not self.flash_timer.isActive():
            self.flash_timer.start()

    def toggle_flash(self):
        self.flash_state = not self.flash_state
        if self.flash_state:
            self.setStyleSheet("""
                QLabel {
                    background-color: #FFFFFF;
                    color: #FF0000;
                    border: 2px solid #800000;
                    padding: 1px 2px;
                    font-family: 'Consolas', 'Courier New', monospace;
                }
            """)
        else:
            self.setStyleSheet("""
                QLabel {
                    background-color: #FF0000;
                    color: #FFFFFF;
                    border: 2px solid #808080;
                    padding: 1px 2px;
                    font-family: 'Consolas', 'Courier New', monospace;
                }
            """)

    def apply_normal_style(self):
        self.setStyleSheet("""
            QLabel {
                background-color: #000000;
                color: #00FF00;
                border: 2px solid #808080;
                padding: 1px 2px;
                font-family: 'Consolas', 'Courier New', monospace;
            }
        """)

    def reset_timer(self, is_net_active=False):
        self.countdown_timer.stop()
        self.flash_timer.stop()
        self.time_remaining = self.INITIAL_TIME_SECONDS
        self.update_display_text()
        self.apply_normal_style()

        if is_net_active:
            self.countdown_timer.start()

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            parent_net_active = False
            main_window = self.window()
            if isinstance(main_window, NetControlApp):
                parent_net_active = main_window.is_net_active

            self.reset_timer(is_net_active=parent_net_active)
        super().mousePressEvent(event)


class NetDurationLabel(QLabel):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.elapsed_seconds = 0
        self.start_time_str = "--:--"
        self.end_time_str = "--:--"

        self.setFrameShape(QFrame.Shape.Box)
        self.setFrameShadow(QFrame.Shadow.Sunken)
        self.setLineWidth(2)
        self.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.setFixedSize(210, 38)

        mono_font = QFont("Consolas", 8, QFont.Weight.Bold)
        mono_font.setStyleHint(QFont.StyleHint.Monospace)
        self.setFont(mono_font)

        self.duration_timer = QTimer(self)
        self.duration_timer.setInterval(1000)
        self.duration_timer.timeout.connect(self.tick)

        self.apply_style()
        self.update_display()

    def apply_style(self):
        self.setStyleSheet("""
            QLabel {
                background-color: #000000;
                color: #00FFFF;
                border: 2px solid #808080;
                padding: 1px 2px;
                font-family: 'Consolas', 'Courier New', monospace;
            }
        """)

    def tick(self):
        self.elapsed_seconds += 1
        self.update_display()

    def update_display(self):
        hrs = self.elapsed_seconds // 3600
        mins = (self.elapsed_seconds % 3600) // 60
        dur_str = f"{hrs:02d}:{mins:02d}"

        line1 = f"DURATION | START {self.start_time_str:>5}"
        line2 = f"   {dur_str:>5} | END   {self.end_time_str:>5}"
        self.setText(f"{line1}\n{line2}")

    def start_duration(self):
        self.start_time_str = datetime.now().strftime("%H:%M")
        self.end_time_str = "--:--"
        if not self.duration_timer.isActive():
            self.duration_timer.start()
        self.update_display()

    def stop_duration(self):
        if self.duration_timer.isActive():
            self.end_time_str = datetime.now().strftime("%H:%M")
        self.duration_timer.stop()
        self.update_display()

    def reset_duration(self):
        self.stop_duration()
        self.elapsed_seconds = 0
        self.start_time_str = "--:--"
        self.end_time_str = "--:--"
        self.update_display()


class ClockBadgeLabel(QLabel):
    def __init__(self, title, is_utc=False, use_24hr=False, parent=None):
        super().__init__(parent)
        self.title = title
        self.is_utc = is_utc
        self.use_24hr = use_24hr

        self.setFrameShape(QFrame.Shape.Box)
        self.setFrameShadow(QFrame.Shadow.Sunken)
        self.setLineWidth(2)
        self.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.setFixedSize(125, 38)

        mono_font = QFont("Consolas", 8, QFont.Weight.Bold)
        mono_font.setStyleHint(QFont.StyleHint.Monospace)
        self.setFont(mono_font)

        self.setStyleSheet("""
            QLabel {
                background-color: #000000;
                color: #00FFFF;
                border: 2px solid #808080;
                padding: 1px 2px;
                font-family: 'Consolas', 'Courier New', monospace;
            }
        """)

        self.timer = QTimer(self)
        self.timer.setInterval(1000)
        self.timer.timeout.connect(self.update_clock)
        self.timer.start()
        self.update_clock()

    def set_24hr_mode(self, use_24hr):
        self.use_24hr = use_24hr
        self.update_clock()

    def update_clock(self):
        if self.is_utc:
            now = datetime.now(timezone.utc)
            fmt = "%H:%M" if self.use_24hr else "%I:%M %p"
        else:
            now = datetime.now()
            fmt = "%H:%M" if self.use_24hr else "%I:%M %p"

        time_str = now.strftime(fmt)
        self.setText(f"{self.title}\n{time_str}")


class ProfileSelectDialog(QDialog):
    def __init__(self, profiles, current_profile="", parent=None):
        super().__init__(parent)
        self.setWindowTitle("Select Net Profile")
        self.setMinimumSize(720, 480)
        self.resize(800, 520)
        self.selected_profile = None

        layout = QVBoxLayout(self)

        lbl = QLabel("Choose Net Control Profile:")
        lbl.setFont(QFont("Tahoma", 9, QFont.Weight.Bold))
        layout.addWidget(lbl)

        modes = [
            ("Regular", "General Nets"),
            ("SKYWARN", "SKYWARN Nets"),
            ("Traffic", "Traffic Nets"),
            ("Training", "Training Nets"),
            ("Resource", "Resource Nets"),
            ("Emergency", "Emergency Nets")
        ]

        self.list_widgets = {}

        scroll_area = QScrollArea()
        scroll_area.setWidgetResizable(True)
        container_widget = QWidget()
        grid_layout = QGridLayout(container_widget)
        grid_layout.setSpacing(10)

        row, col = 0, 0
        for mode_key, mode_title in modes:
            gbox = QGroupBox(mode_title)
            g_layout = QVBoxLayout(gbox)
            g_layout.setContentsMargins(4, 4, 4, 4)

            lwidget = QListWidget()
            self.list_widgets[mode_key] = lwidget
            g_layout.addWidget(lwidget)

            grid_layout.addWidget(gbox, row, col)

            col += 1
            if col >= 2:
                col = 0
                row += 1

        scroll_area.setWidget(container_widget)
        layout.addWidget(scroll_area, 1)

        for name, data in profiles.items():
            mode = data.get("mode", "Regular")
            target_list = self.list_widgets.get(mode, self.list_widgets["Regular"])

            auth_cfg = data.get("auth", {})
            disp_name = f"{name} [LOCKED]" if auth_cfg.get("enabled", False) else name

            item = QListWidgetItem(disp_name)
            item.setData(Qt.ItemDataRole.UserRole, name)
            target_list.addItem(item)
            if name == current_profile:
                target_list.setCurrentItem(item)

        for mode_key, lwidget in self.list_widgets.items():
            lwidget.itemClicked.connect(lambda item, key=mode_key: self.clear_other_selections(key))
            lwidget.itemDoubleClicked.connect(lambda item: self.finish_selection(item.data(Qt.ItemDataRole.UserRole)))

        btn_layout = QHBoxLayout()
        btn_select = EnterClickButton("Select Profile")
        btn_select.clicked.connect(self.accept_selection)

        btn_cancel = QPushButton("Cancel")
        btn_cancel.clicked.connect(self.reject)

        btn_layout.addStretch()
        btn_layout.addWidget(btn_select)
        btn_layout.addWidget(btn_cancel)

        layout.addLayout(btn_layout)

    def clear_other_selections(self, active_key):
        for key, lwidget in self.list_widgets.items():
            if key != active_key:
                lwidget.clearSelection()

    def finish_selection(self, profile_name):
        self.selected_profile = profile_name
        self.accept()

    def accept_selection(self):
        for lwidget in self.list_widgets.values():
            selected = lwidget.selectedItems()
            if selected:
                self.selected_profile = selected[0].data(Qt.ItemDataRole.UserRole)
                self.accept()
                return

class NtsRadiogramDialog(QDialog):
    def __init__(self, default_station="", parent=None):
        super().__init__(parent)
        self.setWindowTitle("ARRL NTS Radiogram Form")
        self.setMinimumSize(680, 700)

        layout = QVBoxLayout(self)

        # Header / Preamble
        preamble_gb = QGroupBox("Message Preamble")
        p_grid = QGridLayout(preamble_gb)

        p_grid.addWidget(QLabel("Number:"), 0, 0)
        self.input_num = QLineEdit()
        p_grid.addWidget(self.input_num, 0, 1)

        p_grid.addWidget(QLabel("Precedence:"), 0, 2)
        self.combo_precedence = QComboBox()
        self.combo_precedence.addItems(["ROUTINE (R)", "WELFARE (W)", "PRIORITY (P)", "EMERGENCY (E)"])
        p_grid.addWidget(self.combo_precedence, 0, 3)

        # Station of Origin
        p_grid.addWidget(QLabel("Station of Origin:"), 1, 0)
        self.input_origin = UppercaseLineEdit()
        self.input_origin.setText(default_station)
        p_grid.addWidget(self.input_origin, 1, 1)

        p_grid.addWidget(QLabel("Check (Word Count):"), 1, 2)
        self.input_check = QLineEdit()
        p_grid.addWidget(self.input_check, 1, 3)

        p_grid.addWidget(QLabel("Place of Origin:"), 1, 4)
        self.input_place = QLineEdit()
        p_grid.addWidget(self.input_place, 1, 5)

        p_grid.addWidget(QLabel("Time Filed:"), 2, 0)
        self.input_time_filed = QLineEdit()
        self.input_time_filed.setText(datetime.now().strftime("%H:%M"))
        p_grid.addWidget(self.input_time_filed, 2, 1)

        p_grid.addWidget(QLabel("Date:"), 2, 2)
        self.input_date = QLineEdit()
        self.input_date.setText(datetime.now().strftime("%b %d"))
        p_grid.addWidget(self.input_date, 2, 3)

        # Handling Instruction (HX) + Variable Input Row
        p_grid.addWidget(QLabel("HX (Handling):"), 0, 4)
        hx_container = QWidget()
        hx_layout = QHBoxLayout(hx_container)
        hx_layout.setContentsMargins(0, 0, 0, 0)
        hx_layout.setSpacing(4)

        self.combo_hx = QComboBox()
        hx_options = [
            ("NONE", "No special handling instructions"),
            ("HXA", "Delivery by email and/or text message preferred"),
            ("HXB", "Cancel message if not delivered within X hours of filing time"),
            ("HXC", "Delivering station report date and time of delivery (TOD) to originating station"),
            ("HXD", "Report identity of station received from/relayed to plus date and time"),
            ("HXE", "Delivering station get reply from addressee, originate message back"),
            ("HXF", "Hold delivery until specified date"),
            ("HXG", "If expense involved, cancel message and service originating station"),
            ("HXI", "Deliver as a radiogram-ICS-213 message"),
            ("HXR", "Confirm addressee has read message and reply to originator"),
            ("HXT", "Special traffic test message")
        ]
        for code, desc in hx_options:
            self.combo_hx.addItem(f"{code} - {desc[:25]}..." if code != "NONE" else "NONE", userData=code)
            self.combo_hx.setItemData(self.combo_hx.count() - 1, desc, Qt.ItemDataRole.ToolTipRole)

        self.input_hx_val = QLineEdit()
        self.input_hx_val.setPlaceholderText("Val (e.g. 48 / Date)")
        self.input_hx_val.setFixedWidth(110)
        self.input_hx_val.setEnabled(False)

        self.combo_hx.currentIndexChanged.connect(self.on_hx_changed)

        hx_layout.addWidget(self.combo_hx, 1)
        hx_layout.addWidget(self.input_hx_val)
        p_grid.addWidget(hx_container, 0, 5)

        layout.addWidget(preamble_gb)

        # Address & Received-By Split Row
        addr_rx_layout = QHBoxLayout()

        # To (Addressee)
        to_gb = QGroupBox("To (Addressee)")
        t_grid = QGridLayout(to_gb)

        t_grid.addWidget(QLabel("Name / Call:"), 0, 0)
        self.input_to_name = QLineEdit()
        t_grid.addWidget(self.input_to_name, 0, 1)

        t_grid.addWidget(QLabel("Phone:"), 1, 0)
        self.input_to_phone = QLineEdit()
        t_grid.addWidget(self.input_to_phone, 1, 1)

        t_grid.addWidget(QLabel("Address:"), 2, 0)
        self.input_to_addr = QLineEdit()
        t_grid.addWidget(self.input_to_addr, 2, 1)

        t_grid.addWidget(QLabel("Email:"), 3, 0)
        self.input_to_email = QLineEdit()
        t_grid.addWidget(self.input_to_email, 3, 1)

        t_grid.addWidget(QLabel("City, State, Zip:"), 4, 0)
        self.input_to_csz = QLineEdit()
        t_grid.addWidget(self.input_to_csz, 4, 1)

        addr_rx_layout.addWidget(to_gb, 1)

        # Optional "Received At" Section
        rx_gb = QGroupBox("This Radio Message Was Received At (Optional)")
        r_grid = QGridLayout(rx_gb)

        r_grid.addWidget(QLabel("Amateur Station:"), 0, 0)
        self.input_rx_station = UppercaseLineEdit()
        r_grid.addWidget(self.input_rx_station, 0, 1)

        r_grid.addWidget(QLabel("Name:"), 1, 0)
        self.input_rx_name = QLineEdit()
        r_grid.addWidget(self.input_rx_name, 1, 1)

        r_grid.addWidget(QLabel("Street:"), 2, 0)
        self.input_rx_street = QLineEdit()
        r_grid.addWidget(self.input_rx_street, 2, 1)

        r_grid.addWidget(QLabel("City, State, Zip:"), 3, 0)
        self.input_rx_csz = QLineEdit()
        r_grid.addWidget(self.input_rx_csz, 3, 1)

        r_grid.addWidget(QLabel("Phone / Email:"), 4, 0)
        self.input_rx_contact = QLineEdit()
        r_grid.addWidget(self.input_rx_contact, 4, 1)

        addr_rx_layout.addWidget(rx_gb, 1)
        layout.addLayout(addr_rx_layout)

        # Message Text
        msg_gb = QGroupBox("Message Text (Word Grouping)")
        m_layout = QVBoxLayout(msg_gb)
        self.text_body = QTextEdit()
        self.text_body.setPlaceholderText("Enter radiogram text here...")
        self.text_body.textChanged.connect(self.auto_calculate_check)
        m_layout.addWidget(self.text_body)
        layout.addWidget(msg_gb, 1)

        # REC'D & SENT Tracking Group
        tracking_gb = QGroupBox("Relay Tracking (REC'D / SENT)")
        track_grid = QGridLayout(tracking_gb)

        track_grid.addWidget(QLabel("REC'D From:"), 0, 0)
        self.input_recd_from = UppercaseLineEdit()
        track_grid.addWidget(self.input_recd_from, 0, 1)

        track_grid.addWidget(QLabel("Date:"), 0, 2)
        self.input_recd_date = QLineEdit()
        track_grid.addWidget(self.input_recd_date, 0, 3)

        track_grid.addWidget(QLabel("Time:"), 0, 4)
        self.input_recd_time = QLineEdit()
        track_grid.addWidget(self.input_recd_time, 0, 5)

        track_grid.addWidget(QLabel("SENT To:"), 1, 0)
        self.input_sent_to = UppercaseLineEdit()
        track_grid.addWidget(self.input_sent_to, 1, 1)

        track_grid.addWidget(QLabel("Date:"), 1, 2)
        self.input_sent_date = QLineEdit()
        track_grid.addWidget(self.input_sent_date, 1, 3)

        track_grid.addWidget(QLabel("Time:"), 1, 4)
        self.input_sent_time = QLineEdit()
        track_grid.addWidget(self.input_sent_time, 1, 5)

        layout.addWidget(tracking_gb)

        # Buttons
        btn_layout = QHBoxLayout()
        btn_save = EnterClickButton("Log & Export PDF (Ctrl+Enter)")
        btn_save.setStyleSheet(BTN_GREEN_STYLE)
        btn_save.clicked.connect(self.accept)

        btn_cancel = QPushButton("Cancel")
        btn_cancel.clicked.connect(self.reject)

        btn_layout.addStretch()
        btn_layout.addWidget(btn_save)
        btn_layout.addWidget(btn_cancel)
        layout.addLayout(btn_layout)

        # Dialog Hotkeys
        self.shortcut_export = QShortcut(QKeySequence("Ctrl+Return"), self)
        self.shortcut_export.activated.connect(self.accept)
        self.shortcut_export_2 = QShortcut(QKeySequence("Ctrl+Enter"), self)
        self.shortcut_export_2.activated.connect(self.accept)

    def on_hx_changed(self):
        code = self.combo_hx.currentData()
        if code in ("HXB", "HXF", "HXT"):
            self.input_hx_val.setEnabled(True)
            if code == "HXB":
                self.input_hx_val.setPlaceholderText("Hours (e.g. 48)")
            elif code == "HXF":
                self.input_hx_val.setPlaceholderText("Date (e.g. SEP 15)")
            elif code == "HXT":
                self.input_hx_val.setPlaceholderText("Test ID#")
            self.input_hx_val.setFocus()
        else:
            self.input_hx_val.clear()
            self.input_hx_val.setEnabled(False)
            self.input_hx_val.setPlaceholderText("Val (Optional)")

    def auto_calculate_check(self):
        text = self.text_body.toPlainText().strip()
        words = [w for w in text.split() if w]
        self.input_check.setText(str(len(words)))

    def get_radiogram_data(self):
        hx_code = self.combo_hx.currentData() or "NONE"
        val = self.input_hx_val.text().strip()

        # Append parameter if code expects a variable
        if hx_code in ("HXB", "HXF", "HXT") and val:
            full_hx = f"{hx_code}{val}"
        else:
            full_hx = hx_code

        return {
            "num": self.input_num.text().strip() or "1",
            "precedence": self.combo_precedence.currentText().split()[0],
            "hx": full_hx,
            "origin": self.input_origin.text().strip().upper() or "",
            "check": self.input_check.text().strip() or "0",
            "place": self.input_place.text().strip() or "-",
            "time_filed": self.input_time_filed.text().strip(),
            "date": self.input_date.text().strip(),
            "to_name": self.input_to_name.text().strip(),
            "to_phone": self.input_to_phone.text().strip(),
            "to_addr": self.input_to_addr.text().strip(),
            "to_email": self.input_to_email.text().strip(),
            "to_csz": self.input_to_csz.text().strip(),
            "rx_station": self.input_rx_station.text().strip().upper(),
            "rx_name": self.input_rx_name.text().strip(),
            "rx_street": self.input_rx_street.text().strip(),
            "rx_csz": self.input_rx_csz.text().strip(),
            "rx_contact": self.input_rx_contact.text().strip(),
            "recd_from": self.input_recd_from.text().strip().upper(),
            "recd_date": self.input_recd_date.text().strip(),
            "recd_time": self.input_recd_time.text().strip(),
            "sent_to": self.input_sent_to.text().strip().upper(),
            "sent_date": self.input_sent_date.text().strip(),
            "sent_time": self.input_sent_time.text().strip(),
            "text": self.text_body.toPlainText().strip()
        }

    def export_radiogram_pdf(self, file_path, data):
        rx_box_html = ""
        if data.get("rx_station") or data.get("rx_name"):
            rx_box_html = f"""
            <td style="width: 50%; vertical-align: top;">
                <b>THIS RADIO MESSAGE WAS RECEIVED AT:</b><br>
                Amateur Station: {data['rx_station']}<br>
                Name: {data['rx_name']}<br>
                Street: {data['rx_street']}<br>
                City, State, Zip: {data['rx_csz']}<br>
                Phone/Email: {data['rx_contact']}
            </td>
            """
        else:
            rx_box_html = """
            <td style="width: 50%; vertical-align: top; color: #808080;">
                <i>[THIS RADIO MESSAGE WAS RECEIVED AT]</i>
            </td>
            """

        html = f"""
        <html>
        <head>
            <style>
                body {{ font-family: Tahoma, sans-serif; margin: 10px; color: #000; }}
                .title {{ text-align: center; font-size: 18pt; font-weight: bold; border-bottom: 2px solid #000; padding-bottom: 5px; }}
                table {{ width: 100%; border-collapse: collapse; margin-top: 8px; }}
                th, td {{ border: 1px solid #000; padding: 4px; font-size: 9.5pt; }}
                th {{ background-color: #E0E0E0; text-align: left; font-size: 8pt; }}
                .body-box {{ border: 1px solid #000; min-height: 140px; padding: 8px; margin-top: 8px; font-family: Courier, monospace; font-size: 11pt; }}
                .track-table td {{ font-size: 8.5pt; height: 35px; vertical-align: top; }}
            </style>
        </head>
        <body>
            <div class="title">ARRL RADIOGRAM</div>
            <table>
                <tr>
                    <th>NUMBER</th><th>PRECEDENCE</th><th>HX</th><th>STATION OF ORIGIN</th><th>CHECK</th><th>PLACE OF ORIGIN</th><th>TIME FILED</th><th>DATE</th>
                </tr>
                <tr>
                    <td>{data['num']}</td><td>{data['precedence']}</td><td>{data['hx']}</td><td>{data['origin']}</td><td>{data['check']}</td><td>{data['place']}</td><td>{data['time_filed']}</td><td>{data['date']}</td>
                </tr>
            </table>
            <table>
                <tr>
                    <td style="width: 50%; vertical-align: top;">
                        <b>TO:</b><br>
                        <b>{data['to_name']}</b><br>
                        {data['to_addr']}<br>
                        {data['to_csz']}<br><br>
                        Phone: {data['to_phone']}<br>
                        Email: {data['to_email']}
                    </td>
                    {rx_box_html}
                </tr>
            </table>
            <div class="body-box">{data['text']}</div>
            
            <table class="track-table">
                <tr>
                    <th style="width: 50%;">REC'D</th>
                    <th style="width: 50%;">SENT</th>
                </tr>
                <tr>
                    <td>
                        FROM: <b>{data['recd_from']}</b><br>
                        DATE: {data['recd_date']} &nbsp;&nbsp;&nbsp;&nbsp; TIME: {data['recd_time']}
                    </td>
                    <td>
                        TO: <b>{data['sent_to']}</b><br>
                        DATE: {data['sent_date']} &nbsp;&nbsp;&nbsp;&nbsp; TIME: {data['sent_time']}
                    </td>
                </tr>
            </table>
        </body>
        </html>
        """
        doc = QTextDocument()
        doc.setHtml(html)
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        printer.setOutputFileName(file_path)
        doc.print(printer)

class ProfileEditorDialog(QDialog):
    def __init__(self, profiles, current_profile="", parent=None):
        super().__init__(parent)
        self.setWindowTitle("Net Profile Creator & Editor")
        self.resize(680, 840)
        self.profiles = json.loads(json.dumps(profiles))
        self.active_profile_name = None

        layout = QVBoxLayout(self)

        top_row = QHBoxLayout()
        top_row.addWidget(QLabel("Select Profile to Edit:"))
        self.combo_profiles = QComboBox()
        self.combo_profiles.currentIndexChanged.connect(self.load_selected_profile)
        top_row.addWidget(self.combo_profiles, 1)

        btn_new_prof = QPushButton("New Profile")
        btn_new_prof.clicked.connect(self.create_new_profile)
        btn_del_prof = QPushButton("Delete")
        btn_del_prof.clicked.connect(self.delete_current_profile)

        top_row.addWidget(btn_new_prof)
        top_row.addWidget(btn_del_prof)
        layout.addLayout(top_row)

        mode_row = QHBoxLayout()
        mode_row.addWidget(QLabel("Profile Name:"))
        self.input_prof_name = QLineEdit()
        mode_row.addWidget(self.input_prof_name, 1)

        mode_row.addWidget(QLabel("Net Mode:"))
        self.combo_mode = QComboBox()
        self.combo_mode.addItems(["Regular", "SKYWARN", "Traffic", "Training", "Resource", "Emergency"])
        self.combo_mode.currentTextChanged.connect(self.on_mode_changed)
        mode_row.addWidget(self.combo_mode)
        layout.addLayout(mode_row)

        options_row = QHBoxLayout()
        self.chk_has_roll_call = EnterToggleCheckBox("Roll Call")
        self.chk_has_nts = EnterToggleCheckBox("Radiogram")
        self.chk_custom_statuses = EnterToggleCheckBox("Custom Status Badging")
        self.chk_use_24hr_clock = EnterToggleCheckBox("24-Hour Clock")

        self.btn_repeater_info = QPushButton("Repeater Info")
        self.btn_repeater_info.clicked.connect(self.open_repeater_info_dialog)

        options_row.addWidget(self.chk_has_roll_call)
        options_row.addSpacing(12)
        options_row.addWidget(self.chk_custom_statuses)
        options_row.addSpacing(12)
        options_row.addWidget(self.chk_use_24hr_clock)
        options_row.addSpacing(12)
        options_row.addWidget(self.chk_has_nts)
        options_row.addSpacing(12)
        options_row.addWidget(self.btn_repeater_info)
        options_row.addStretch()
        layout.addLayout(options_row)

        self.auth_group = QGroupBox("Access Control Security (Optional Passcode Protection)")
        auth_layout = QHBoxLayout(self.auth_group)

        self.chk_enable_auth = EnterToggleCheckBox("Require Passcode / PIN to Open Profile")
        self.chk_enable_auth.toggled.connect(self.toggle_auth_btn)
        auth_layout.addWidget(self.chk_enable_auth)

        self.btn_user_mgmt = QPushButton("User Management")
        self.btn_user_mgmt.clicked.connect(self.open_user_management)
        auth_layout.addWidget(self.btn_user_mgmt)
        auth_layout.addStretch()

        layout.addWidget(self.auth_group)

        types_row = QHBoxLayout()

        font_metrics = QFontMetrics(self.font())
        line_height = font_metrics.lineSpacing()
        five_lines_height = line_height * 5 + 12

        station_group = QGroupBox("Selectable Station Types (One per line)")
        st_layout = QVBoxLayout(station_group)
        self.text_station_types = QTextEdit()
        self.text_station_types.setFixedHeight(five_lines_height)
        st_layout.addWidget(self.text_station_types)
        types_row.addWidget(station_group, 1)

        traffic_group = QGroupBox("Traffic / Condition Types (One per line)")
        tr_layout = QVBoxLayout(traffic_group)
        self.text_traffic_types = QTextEdit()
        self.text_traffic_types.setFixedHeight(five_lines_height)
        tr_layout.addWidget(self.text_traffic_types)
        types_row.addWidget(traffic_group, 1)

        layout.addLayout(types_row)

        self.statuses_group = QGroupBox("Net Statuses Color Map (Hex Code e.g. #FF0000 or #FFFF00)")
        statuses_layout = QVBoxLayout(self.statuses_group)

        self.table_statuses = QTableWidget(0, 3)
        self.table_statuses.setHorizontalHeaderLabels(["Status Name", "Text Hex", "Background Hex"])
        self.table_statuses.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Interactive)
        self.table_statuses.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        self.table_statuses.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
        self.table_statuses.setColumnWidth(0, 160)

        header_height = self.table_statuses.horizontalHeader().height() or 25
        row_height = 24
        table_four_rows_height = header_height + (row_height * 4) + 6
        self.table_statuses.setFixedHeight(table_four_rows_height)

        st_btns = QHBoxLayout()
        btn_add_st = QPushButton("Add Status")
        btn_add_st.clicked.connect(self.add_status_row)
        btn_del_st = QPushButton("Remove Status")
        btn_del_st.clicked.connect(self.delete_status_row)
        st_btns.addWidget(btn_add_st)
        st_btns.addWidget(btn_del_st)
        st_btns.addStretch()

        statuses_layout.addWidget(self.table_statuses)
        statuses_layout.addLayout(st_btns)
        layout.addWidget(self.statuses_group)

        self.api_group = QGroupBox("Web Status API Sync Dropdown Options")
        api_layout = QHBoxLayout(self.api_group)

        self.chk_enable_api = EnterToggleCheckBox("Enable Online Status API Sync for this Profile")
        self.chk_enable_api.toggled.connect(self.toggle_api_btn)
        api_layout.addWidget(self.chk_enable_api)

        self.btn_config_api = QPushButton("Configure")
        self.btn_config_api.clicked.connect(self.open_api_config)
        api_layout.addWidget(self.btn_config_api)
        api_layout.addStretch()

        layout.addWidget(self.api_group)

        self.scripts_group = QGroupBox("Script Tabs Management (Press Shift+Enter for line breaks)")
        self.scripts_group.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        scripts_layout = QVBoxLayout(self.scripts_group)

        self.table_scripts = QTableWidget(0, 2)
        self.table_scripts.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self.table_scripts.setHorizontalHeaderLabels(["Tab Title", "Script Preamble Text"])
        self.table_scripts.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Interactive)
        self.table_scripts.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        self.table_scripts.setColumnWidth(0, 140)

        sc_btns = QHBoxLayout()
        btn_add_script = QPushButton("Add Tab")
        btn_add_script.clicked.connect(self.add_script_row)
        btn_del_script = QPushButton("Remove Selected Tab")
        btn_del_script.clicked.connect(self.delete_script_row)

        sc_btns.addWidget(btn_add_script)
        sc_btns.addWidget(btn_del_script)
        sc_btns.addStretch()

        scripts_layout.addWidget(self.table_scripts, 1)
        scripts_layout.addLayout(sc_btns)
        layout.addWidget(self.scripts_group, 1)

        bottom_btns = QHBoxLayout()

        btn_save_single = QPushButton("Save Profile")
        btn_save_single.setStyleSheet(BTN_SAVE_STYLE)
        btn_save_single.clicked.connect(self.save_single_profile)

        btn_save_all = EnterClickButton("Save All Profiles Close")
        btn_save_all.clicked.connect(self.save_and_close)

        btn_cancel = QPushButton("Cancel")
        btn_cancel.clicked.connect(self.reject)

        bottom_btns.addWidget(btn_save_single)
        bottom_btns.addStretch()
        bottom_btns.addWidget(btn_save_all)
        bottom_btns.addWidget(btn_cancel)
        layout.addLayout(bottom_btns)

        self.temp_repeater_data = {"freq": "UNSET", "offset": "X", "tone": "UNSET"}
        self.temp_auth_users = []
        self.temp_api_sync = {}
        self.temp_web_opts = []

        self.populate_combo(current_profile)

    def open_repeater_info_dialog(self):
        dialog = RepeaterInfoDialog(self.temp_repeater_data, parent=self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.temp_repeater_data = dialog.get_repeater_data()

    def toggle_api_btn(self, enabled):
        self.btn_config_api.setVisible(enabled)

    def toggle_auth_btn(self, enabled):
        self.btn_user_mgmt.setVisible(enabled)

    def open_api_config(self):
        mode_text = self.combo_mode.currentText()
        dialog = ApiConfigDialog(self.temp_api_sync, self.temp_web_opts, mode_text=mode_text, parent=self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            updated_api, updated_opts = dialog.get_config_data()
            self.temp_api_sync = updated_api
            self.temp_web_opts = updated_opts

    def open_user_management(self):
        dialog = UserManagementDialog(self.temp_auth_users, parent=self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.temp_auth_users = dialog.get_users_data()

    def on_mode_changed(self, mode_text):
        if mode_text == "SKYWARN":
            self.chk_has_roll_call.setEnabled(False)
            self.chk_has_roll_call.setChecked(False)
            self.chk_custom_statuses.setChecked(True)
        else:
            self.chk_has_roll_call.setEnabled(True)

    def populate_combo(self, select_name=""):
        self.combo_profiles.blockSignals(True)
        self.combo_profiles.clear()

        categories = {
            "Regular": [],
            "SKYWARN": [],
            "Traffic": [],
            "Training": [],
            "Resource": [],
            "Emergency": []
        }

        for p_name, p_data in self.profiles.items():
            mode = p_data.get("mode", "Regular")
            if mode in categories:
                categories[mode].append(p_name)
            else:
                categories["Regular"].append(p_name)

        target_idx = -1

        for cat_name, p_list in categories.items():
            if p_list:
                p_list.sort()
                header_text = f"--- {cat_name.upper()} NETS ---"
                self.combo_profiles.addItem(header_text)

                header_idx = self.combo_profiles.count() - 1
                item = self.combo_profiles.model().item(header_idx)
                if item:
                    item.setEnabled(False)

                for prof_name in p_list:
                    self.combo_profiles.addItem(prof_name)
                    if prof_name == select_name:
                        target_idx = self.combo_profiles.count() - 1

        self.combo_profiles.blockSignals(False)

        if target_idx >= 0:
            self.combo_profiles.setCurrentIndex(target_idx)
        else:
            for i in range(self.combo_profiles.count()):
                item = self.combo_profiles.model().item(i)
                if item and item.isEnabled():
                    self.combo_profiles.setCurrentIndex(i)
                    break

        self.load_selected_profile()

    def load_selected_profile(self):
        if self.active_profile_name and self.active_profile_name in self.profiles:
            self.sync_form_to_data(self.active_profile_name)

        name = self.combo_profiles.currentText()
        if not name or name.startswith("--- ") or name not in self.profiles:
            return

        self.active_profile_name = name
        prof = self.profiles[name]

        self.input_prof_name.setText(name)
        mode = prof.get("mode", "Regular")
        self.combo_mode.setCurrentText(mode)
        self.chk_has_roll_call.setChecked(prof.get("has_roll_call", False))
        self.chk_custom_statuses.setChecked(prof.get("use_custom_statuses", mode == "SKYWARN"))
        self.chk_use_24hr_clock.setChecked(prof.get("use_24hr_clock", False))
        self.chk_has_nts.setChecked(prof.get("has_nts", False))
        self.on_mode_changed(mode)

        self.temp_repeater_data = prof.get("repeater", {"freq": "UNSET", "offset": "X", "tone": "UNSET"})

        auth_cfg = prof.get("auth", {})
        is_auth_enabled = auth_cfg.get("enabled", False)
        self.chk_enable_auth.setChecked(is_auth_enabled)

        if "users" in auth_cfg:
            self.temp_auth_users = auth_cfg.get("users", [])
        else:
            passc = auth_cfg.get("passcode", "")
            calls = auth_cfg.get("allowed_callsigns", [])
            self.temp_auth_users = [{"callsign": c, "passcode": passc} for c in calls] if calls else [{"callsign": "W9A", "passcode": passc}]

        self.toggle_auth_btn(is_auth_enabled)

        self.text_station_types.setText("\n".join(prof.get("station_types", [])))
        self.text_traffic_types.setText("\n".join(prof.get("traffic_types", [])))
        self.temp_web_opts = prof.get("web_status_options", ["ACTIVE", "STANDBY", "CLOSED"])

        net_statuses = prof.get("net_statuses", {})
        self.table_statuses.setRowCount(0)
        for st_name, val in net_statuses.items():
            row = self.table_statuses.rowCount()
            self.table_statuses.insertRow(row)
            self.table_statuses.setItem(row, 0, QTableWidgetItem(st_name))

            t_hex, b_hex = parse_color_map(val)
            self.table_statuses.setItem(row, 1, QTableWidgetItem(t_hex))
            self.table_statuses.setItem(row, 2, QTableWidgetItem(b_hex))

        self.temp_api_sync = prof.get("api_sync", {
            "enabled": False,
            "url": "https://www.mcinares.org/chatgemilaude/doupdate/",
            "qrz": "W9A",
            "param_name": "mcinaresstatus",
            "active_value": "ACTIVE",
            "inactive_value": "INACTIVE"
        })
        is_api_enabled = self.temp_api_sync.get("enabled", False)
        self.chk_enable_api.setChecked(is_api_enabled)
        self.toggle_api_btn(is_api_enabled)

        scripts = prof.get("scripts", [])
        self.table_scripts.setRowCount(0)
        for row, sc in enumerate(scripts):
            self.table_scripts.insertRow(row)
            self.table_scripts.setItem(row, 0, QTableWidgetItem(sc.get("title", "")))

            editor = ScriptTextEdit()
            editor.setPlainText(sc.get("text", "").replace("\\n", "\n"))
            self.table_scripts.setCellWidget(row, 1, editor)
            self.table_scripts.setRowHeight(row, 75)

    def sync_form_to_data(self, target_name):
        if target_name not in self.profiles:
            return

        st_types = [t.strip() for t in self.text_station_types.toPlainText().split("\n") if t.strip()]
        tr_types = [t.strip() for t in self.text_traffic_types.toPlainText().split("\n") if t.strip()]

        statuses_map = {}
        for r in range(self.table_statuses.rowCount()):
            s_item = self.table_statuses.item(r, 0)
            t_item = self.table_statuses.item(r, 1)
            b_item = self.table_statuses.item(r, 2)

            s_name = s_item.text().strip().upper() if s_item else ""
            t_hex = t_item.text().strip() if t_item else "#FFFF00"
            b_hex = b_item.text().strip() if b_item else "#000000"

            if s_name:
                statuses_map[s_name] = {"text": t_hex, "bg": b_hex}

        if not statuses_map:
            statuses_map = {
                "ACTIVE": {"text": "#000000", "bg": "#00FF00"},
                "CLOSED": {"text": "#00FF00", "bg": "#000000"},
                "STANDBY": {"text": "#FFFF00", "bg": "#000000"}
            }

        scripts = []
        for r in range(self.table_scripts.rowCount()):
            t_item = self.table_scripts.item(r, 0)
            c_widget = self.table_scripts.cellWidget(r, 1)
            title = t_item.text().strip() if t_item else "Script"
            text = c_widget.toPlainText() if isinstance(c_widget, QTextEdit) else ""
            if title:
                scripts.append({"title": title, "text": text})

        self.temp_api_sync["enabled"] = self.chk_enable_api.isChecked()

        self.profiles[target_name] = {
            "mode": self.combo_mode.currentText(),
            "has_roll_call": self.chk_has_roll_call.isChecked(),
            "use_custom_statuses": self.chk_custom_statuses.isChecked(),
            "use_24hr_clock": self.chk_use_24hr_clock.isChecked(),
            "repeater": self.temp_repeater_data,
            "auth": {
                "enabled": self.chk_enable_auth.isChecked(),
                "users": self.temp_auth_users
            },
            "station_types": st_types,
            "traffic_types": tr_types,
            "web_status_options": self.temp_web_opts,
            "net_statuses": statuses_map,
            "scripts": scripts,
            "api_sync": self.temp_api_sync
        }
        self.profiles[target_name]["has_nts"] = self.chk_has_nts.isChecked()

    def create_new_profile(self):
        name, ok = QInputDialog.getText(self, "New Profile", "Enter Profile Name:")
        if ok and name.strip():
            name = name.strip()
            if name in self.profiles:
                QMessageBox.warning(self, "Duplicate Profile", f"Profile '{name}' already exists.")
                return
            self.profiles[name] = {
                "mode": "Regular",
                "has_roll_call": True,
                "use_custom_statuses": False,
                "use_24hr_clock": False,
                "repeater": {
                    "freq": "146.700",
                    "offset": "-",
                    "tone": "88.5 PL"
                },
                "auth": {
                    "enabled": False,
                    "users": [{"callsign": "W9A", "passcode": "1234"}]
                },
                "station_types": ["Base", "Mobile", "NCS"],
                "traffic_types": ["Question", "Announcement", "Priority", "Emergency", "General Note"],
                "web_status_options": ["ACTIVE", "STANDBY", "CLOSED"],
                "net_statuses": {
                    "ACTIVE": {"text": "#000000", "bg": "#00FF00"},
                    "CLOSED": {"text": "#00FF00", "bg": "#000000"},
                    "STANDBY": {"text": "#FFFF00", "bg": "#000000"}
                },
                "scripts": [{"title": "Preamble", "text": "Welcome to the net..."}],
                "api_sync": {
                    "enabled": False,
                    "url": "https://www.mcinares.org/chatgemilaude/doupdate/",
                    "qrz": "W9A",
                    "param_name": "mcinaresstatus",
                    "active_value": "ACTIVE",
                    "inactive_value": "INACTIVE"
                }
            }
            self.populate_combo(name)

    def delete_current_profile(self):
        curr = self.combo_profiles.currentText()
        if not curr or curr.startswith("--- ") or curr not in self.profiles:
            return

        if len(self.profiles) <= 1:
            QMessageBox.warning(self, "Cannot Delete", "You must keep at least one profile.")
            return

        dialog = RedDangerDialog("Confirm Delete Profile", f"Delete profile '{curr}'?", require_type_delete=True, parent=self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            del self.profiles[curr]
            self.active_profile_name = None
            self.populate_combo()

    def add_status_row(self):
        row = self.table_statuses.rowCount()
        self.table_statuses.insertRow(row)
        self.table_statuses.setItem(row, 0, QTableWidgetItem("NEW_STATUS"))
        self.table_statuses.setItem(row, 1, QTableWidgetItem("#FFFF00"))
        self.table_statuses.setItem(row, 2, QTableWidgetItem("#000000"))

    def delete_status_row(self):
        row = self.table_statuses.currentRow()
        if row >= 0:
            self.table_statuses.removeRow(row)

    def add_script_row(self):
        row = self.table_scripts.rowCount()
        self.table_scripts.insertRow(row)
        self.table_scripts.setItem(row, 0, QTableWidgetItem(f"Script {row + 1}"))

        editor = ScriptTextEdit()
        self.table_scripts.setCellWidget(row, 1, editor)
        self.table_scripts.setRowHeight(row, 75)

    def delete_script_row(self):
        row = self.table_scripts.currentRow()
        if row >= 0:
            self.table_scripts.removeRow(row)

    def save_single_profile(self):
        if self.active_profile_name:
            new_name = self.input_prof_name.text().strip()
            if not new_name:
                QMessageBox.warning(self, "Validation Error", "Profile name cannot be empty.")
                return

            if new_name != self.active_profile_name:
                self.profiles[new_name] = self.profiles.pop(self.active_profile_name)
                self.active_profile_name = new_name

            self.sync_form_to_data(new_name)
            save_profiles_to_disk(self.profiles)
            self.populate_combo(new_name)
            QMessageBox.information(self, "Profile Saved", f"Profile '{new_name}' has been saved successfully.")

    def save_and_close(self):
        if self.active_profile_name:
            new_name = self.input_prof_name.text().strip()
            if not new_name:
                QMessageBox.warning(self, "Validation Error", "Profile name cannot be empty.")
                return

            if new_name != self.active_profile_name:
                self.profiles[new_name] = self.profiles.pop(self.active_profile_name)
                self.active_profile_name = new_name

            self.sync_form_to_data(new_name)

        save_profiles_to_disk(self.profiles)
        self.accept()


def parse_roll_call_file(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb.active
    roster = []

    column_blocks = [(2, 3, 4), (7, 8, 9)]

    for p_col, s_col, n_col in column_blocks:
        for row in range(10, sheet.max_row + 1):
            prefix = sheet.cell(row=row, column=p_col).value
            suffix = sheet.cell(row=row, column=s_col).value
            name = sheet.cell(row=row, column=n_col).value

            if prefix and suffix:
                callsign = f"{str(prefix).strip()}{str(suffix).strip()}".upper()
                op_name = str(name).strip() if name else ""
                roster.append({"callsign": callsign, "name": op_name, "prefix": str(prefix).strip(), "suffix": str(suffix).strip()})

    return roster


def export_xlsx_checkin_list(stations, output_path):
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Net Check-Ins"

    sheet.merge_cells("A1:B1")
    sheet["A1"] = f"Net Check-In List — {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    sheet["A1"].font = openpyxl.styles.Font(size=14, bold=True)

    sheet.cell(row=3, column=1, value="Callsign Prefix / Suffix")
    sheet.cell(row=3, column=2, value="Operator Name / Details")

    for idx, st in enumerate(stations, start=4):
        call = st.get("callsign", "")
        name = st.get("info", st.get("detail", ""))

        sheet.cell(row=idx, column=1, value=call)
        sheet.cell(row=idx, column=2, value=name)

    wb.save(output_path)


class RollCallDialog(QDialog):
    def __init__(self, roster, start_index=0, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Net Roll Call")
        self.setFixedSize(400, 240)
        self.roster = roster
        self.index = start_index
        self.main_app = parent

        layout = QVBoxLayout(self)
        layout.setContentsMargins(8, 8, 8, 8)
        layout.setSpacing(6)

        self.lbl_progress = QLabel()
        self.lbl_progress.setStyleSheet("color: #000080; font-weight: bold;")

        self.lbl_callsign = QLabel()
        self.lbl_callsign.setFont(QFont("Consolas", 22, QFont.Weight.Bold))
        self.lbl_callsign.setAlignment(Qt.AlignmentFlag.AlignCenter)

        self.lbl_name = QLabel()
        self.lbl_name.setFont(QFont("Tahoma", 11))
        self.lbl_name.setAlignment(Qt.AlignmentFlag.AlignCenter)

        opt_layout = QHBoxLayout()
        opt_layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.chk_traffic = EnterToggleCheckBox("Traffic")
        self.chk_announcement = EnterToggleCheckBox("Announcement")
        opt_layout.addWidget(self.chk_traffic)
        opt_layout.addSpacing(15)
        opt_layout.addWidget(self.chk_announcement)

        btn_layout = QHBoxLayout()
        self.btn_prev = QPushButton("Previous\n(Ctrl+Up)")
        self.btn_prev.clicked.connect(self.prev_station)

        self.btn_checkin = EnterClickButton("Check In\n(Ctrl+Space)")
        self.btn_checkin.setStyleSheet(BTN_GREEN_STYLE)
        self.btn_checkin.clicked.connect(self.check_in_current)

        self.btn_skip = QPushButton("Skip\n(Ctrl+Down)")
        self.btn_skip.clicked.connect(self.next_station)

        btn_layout.addWidget(self.btn_prev)
        btn_layout.addWidget(self.btn_checkin)
        btn_layout.addWidget(self.btn_skip)

        layout.addWidget(self.lbl_progress)
        layout.addWidget(self.lbl_callsign)
        layout.addWidget(self.lbl_name)
        layout.addLayout(opt_layout)
        layout.addLayout(btn_layout)

        self.shortcut_checkin = QShortcut(QKeySequence("Ctrl+Space"), self)
        self.shortcut_checkin.activated.connect(self.check_in_current)

        self.shortcut_skip = QShortcut(QKeySequence("Ctrl+Down"), self)
        self.shortcut_skip.activated.connect(self.next_station)

        self.shortcut_prev = QShortcut(QKeySequence("Ctrl+Up"), self)
        self.shortcut_prev.activated.connect(self.prev_station)

        self.seek_next_valid(forward=True)

    def is_already_checked_in(self, callsign):
        if self.main_app and isinstance(self.main_app, NetControlApp):
            return any(st["callsign"] == callsign for st in self.main_app.stations)
        return False

    def seek_next_valid(self, forward=True):
        while 0 <= self.index < len(self.roster):
            call = self.roster[self.index]["callsign"]
            if self.is_already_checked_in(call):
                self.index += (1 if forward else -1)
            else:
                break
        self.update_display()

    def update_display(self):
        if 0 <= self.index < len(self.roster):
            st = self.roster[self.index]
            self.lbl_progress.setText(f"Station {self.index + 1} of {len(self.roster)}")
            self.lbl_callsign.setText(st["callsign"])
            self.lbl_name.setText(f"Operator: {st['name']}")
            self.chk_traffic.setChecked(False)
            self.chk_announcement.setChecked(False)
        elif self.index >= len(self.roster):
            self.accept()

    def check_in_current(self):
        if 0 <= self.index < len(self.roster):
            st = self.roster[self.index]
            has_traffic = self.chk_traffic.isChecked()
            has_ann = self.chk_announcement.isChecked()

            items = []
            if has_traffic:
                items.append("Traffic")
            if has_ann:
                items.append("Announcement")
            col_detail = ", ".join(items) if items else "-"

            if self.main_app and isinstance(self.main_app, NetControlApp):
                self.main_app.add_rollcall_checkin(st["callsign"], col_detail, st["name"])

            self.next_station()

    def next_station(self):
        self.index += 1
        self.seek_next_valid(forward=True)

    def prev_station(self):
        if self.index > 0:
            self.index -= 1
            self.seek_next_valid(forward=False)


class NetControlApp(QMainWindow):
    AUTOSAVE_FILE = os.path.join(BASE_DIR, "autosave_session.json")

    def __init__(self):
        super().__init__()
        self.setWindowTitle("Amateur Radio Net Control Manager")
        self.setMinimumSize(1280, 1024)
        self.resize(1280, 1024)

        self.base_font_size = 10

        self.is_net_active = False
        self.net_state_name = "INACTIVE"
        self.current_ncs_callsign = "UNASSIGNED"
        self.current_scribe_callsign = ""
        self.current_mode = "Regular"
        self.has_roll_call = False
        self.use_custom_statuses = False
        self.use_24hr_clock = False
        self.net_statuses = {}
        self.stations = []
        self.activities = []
        
        self.rollcall_roster = []
        self.rollcall_index = 0

        self.profiles = load_profiles_from_disk()
        self.current_profile_name = ""

        self.station_memory = load_station_memory()

        self.init_ui()
        self.apply_font_size()
        self.setup_shortcuts()
        self.setup_autosave()
        self.setup_api_polling()

    def setup_api_polling(self):
        self.api_poll_timer = QTimer(self)
        self.api_poll_timer.setInterval(30000)
        self.api_poll_timer.timeout.connect(self.poll_remote_web_status)
        self.api_poll_timer.start()

    def poll_remote_web_status(self):
        prof = self.profiles.get(self.current_profile_name, {})
        api_cfg = prof.get("api_sync", {})
        if not api_cfg or not api_cfg.get("enabled", False):
            self.web_status_badge.set_status("SYNC OFF")
            return

        def worker():
            update_url = api_cfg.get("url", "").strip()
            param = api_cfg.get("param_name", "").strip()
            if not param:
                param = "skywarnstatus" if self.current_mode == "SKYWARN" else "mcinaresstatus"

            if not update_url:
                return

            read_url = update_url.rstrip('/')
            if read_url.endswith("doupdate"):
                read_url = read_url[:-8] + "getstatus"
            elif not read_url.endswith("getstatus"):
                read_url += "/getstatus"

            try:
                req = urllib.request.Request(read_url, method="GET")
                req.add_header("User-Agent", "NCM4-DesktopApp")

                with urllib.request.urlopen(req, timeout=5) as resp:
                    raw_data = resp.read().decode("utf-8", errors="ignore").strip()

                    remote_val = ""
                    try:
                        data = json.loads(raw_data)
                        if isinstance(data, dict) and param in data:
                            remote_val = str(data.get(param, "")).strip()
                        elif isinstance(data, str):
                            remote_val = data.strip()
                    except Exception:
                        cleaned = re.sub(r'<[^>]+>', '', raw_data).strip()
                        cleaned = " ".join(cleaned.split())
                        if not cleaned.startswith("YOU SAID"):
                            remote_val = cleaned

                    if remote_val and "404" not in remote_val:
                        self.web_status_badge.set_status(remote_val)

            except Exception as e:
                print(f"[API Read Error]: {e}")

        threading.Thread(target=worker, daemon=True).start()

    def apply_font_size(self):
        app = QApplication.instance()
        if app:
            app.setStyleSheet(generate_stylesheet(self.base_font_size))

    def increase_font_size(self):
        if self.base_font_size < 20:
            self.base_font_size += 2
            self.apply_font_size()

    def decrease_font_size(self):
        if self.base_font_size > 8:
            self.base_font_size -= 2
            self.apply_font_size()

    def open_hotkey_help(self):
        dialog = HotkeyHelpDialog(self)
        dialog.exec()

    def setup_shortcuts(self):
        # --- Global & Navigation Shortcuts ---
        self.shortcut_toggle_legends = QShortcut(QKeySequence("F1"), self)
        self.shortcut_toggle_legends.activated.connect(self.open_hotkey_help)

        self.shortcut_select_profile = QShortcut(QKeySequence("F2"), self)
        self.shortcut_select_profile.activated.connect(self.open_profile_select_dialog)

        self.shortcut_edit_profile = QShortcut(QKeySequence("F3"), self)
        self.shortcut_edit_profile.activated.connect(self.open_profile_editor_dialog)

        self.shortcut_toggle_net = QShortcut(QKeySequence("F4"), self)
        self.shortcut_toggle_net.activated.connect(self.toggle_net_mode)

        self.shortcut_transfer_net = QShortcut(QKeySequence("F5"), self)
        self.shortcut_transfer_net.activated.connect(self.transfer_net_control)

        self.shortcut_save_session = QShortcut(QKeySequence("Ctrl+S"), self)
        self.shortcut_save_session.activated.connect(self.manual_save_session)

        self.shortcut_open_session = QShortcut(QKeySequence("Ctrl+O"), self)
        self.shortcut_open_session.activated.connect(self.open_saved_session_file)

        self.shortcut_reset_net = QShortcut(QKeySequence("Ctrl+R"), self)
        self.shortcut_reset_net.activated.connect(self.reset_entire_net)

        self.shortcut_zoom_in_1 = QShortcut(QKeySequence("Ctrl++"), self)
        self.shortcut_zoom_in_1.activated.connect(self.increase_font_size)
        self.shortcut_zoom_in_2 = QShortcut(QKeySequence("Ctrl+="), self)
        self.shortcut_zoom_in_2.activated.connect(self.increase_font_size)

        self.shortcut_zoom_out_1 = QShortcut(QKeySequence("Ctrl+-"), self)
        self.shortcut_zoom_out_1.activated.connect(self.decrease_font_size)
        self.shortcut_zoom_out_2 = QShortcut(QKeySequence("Ctrl+_"), self)
        self.shortcut_zoom_out_2.activated.connect(self.decrease_font_size)

        # --- Station Management (Check-Ins & Roster) ---
        self.shortcut_focus_callsign = QShortcut(QKeySequence("Alt+1"), self)
        self.shortcut_focus_callsign.activated.connect(self.focus_callsign_input)

        self.shortcut_roll_call = QShortcut(QKeySequence("F8"), self)
        self.shortcut_roll_call.activated.connect(self.start_roll_call)

        self.shortcut_search_roster = QShortcut(QKeySequence("Ctrl+F"), self)
        self.shortcut_search_roster.activated.connect(self.open_combined_station_search_edit)

        self.shortcut_standby = QShortcut(QKeySequence("Ctrl+T"), self)
        self.shortcut_standby.activated.connect(self.toggle_station_standby)

        self.shortcut_checkout = QShortcut(QKeySequence("Ctrl+U"), self)
        self.shortcut_checkout.activated.connect(self.checkout_station)

        self.shortcut_delete_station = QShortcut(QKeySequence("Ctrl+Shift+Delete"), self)
        self.shortcut_delete_station.activated.connect(self.delete_station)

        # --- Traffic & Activity Logging ---
        self.shortcut_focus_traffic_call = QShortcut(QKeySequence("Alt+2"), self)
        self.shortcut_focus_traffic_call.activated.connect(self.focus_traffic_station_input)

        self.shortcut_log_traffic_1 = QShortcut(QKeySequence("Ctrl+Return"), self)
        self.shortcut_log_traffic_1.activated.connect(self.log_traffic_item)
        self.shortcut_log_traffic_2 = QShortcut(QKeySequence("Ctrl+Enter"), self)
        self.shortcut_log_traffic_2.activated.connect(self.log_traffic_item)

        self.shortcut_nts_dialog = QShortcut(QKeySequence("Ctrl+Shift+T"), self)
        self.shortcut_nts_dialog.activated.connect(self.open_nts_dialog)

        self.shortcut_traffic_now = QShortcut(QKeySequence("Ctrl+N"), self)
        self.shortcut_traffic_now.activated.connect(self.set_traffic_time_now)

        self.shortcut_clear_traffic_f12 = QShortcut(QKeySequence("F12"), self)
        self.shortcut_clear_traffic_f12.activated.connect(self.clear_traffic_form)
        self.shortcut_clear_traffic_esc = QShortcut(QKeySequence("Esc"), self)
        self.shortcut_clear_traffic_esc.activated.connect(self.clear_traffic_form)

        self.shortcut_traffic_search = QShortcut(QKeySequence("Ctrl+Shift+F"), self)
        self.shortcut_traffic_search.activated.connect(self.open_combined_traffic_search_edit)

        self.shortcut_traffic_delete = QShortcut(QKeySequence("Ctrl+Delete"), self)
        self.shortcut_traffic_delete.activated.connect(self.delete_traffic_entry)

        # --- Scripts, Web API & Badges ---
        self.shortcut_next_script = QShortcut(QKeySequence("Ctrl+]"), self)
        self.shortcut_next_script.activated.connect(self.next_script_tab)

        self.shortcut_prev_script = QShortcut(QKeySequence("Ctrl+["), self)
        self.shortcut_prev_script.activated.connect(self.prev_script_tab)

        self.shortcut_focus_web_status = QShortcut(QKeySequence("F9"), self)
        self.shortcut_focus_web_status.activated.connect(self.focus_web_status_dropdown)

        self.shortcut_send_web_status = QShortcut(QKeySequence("F10"), self)
        self.shortcut_send_web_status.activated.connect(self.push_custom_web_status)

        self.shortcut_set_status = QShortcut(QKeySequence("Ctrl+Alt+S"), self)
        self.shortcut_set_status.activated.connect(self.apply_selected_net_status)

    def setup_autosave(self):
        self.autosave_timer = QTimer(self)
        self.autosave_timer.setInterval(15000)
        self.autosave_timer.timeout.connect(self.autosave_session)
        self.autosave_timer.start()

    def init_ui(self):
        main_widget = QWidget()
        self.setCentralWidget(main_widget)

        main_layout = QVBoxLayout(main_widget)
        main_layout.setSizeConstraint(QLayout.SizeConstraint.SetNoConstraint)
        main_layout.setContentsMargins(4, 4, 4, 4)
        main_layout.setSpacing(3)

        # Header Layout
        header_layout = QHBoxLayout()
        header_layout.setContentsMargins(2, 2, 2, 2)
        header_layout.setSpacing(0)

        # 1. Left Subgroup (Profile & Web Status)
        left_subgroup = QHBoxLayout()
        left_subgroup.setContentsMargins(0, 0, 0, 0)
        left_subgroup.setSpacing(6)
        
        self.profile_badge = ProfileBadgeLabel()
        self.web_status_badge = WebApiStatusBadgeLabel()
        
        left_subgroup.addWidget(self.profile_badge)
        left_subgroup.addWidget(self.web_status_badge)

        # 2. Center Subgroup (Start Net, Status, ID Timer)
        self.center_group = QWidget()
        self.center_layout = QHBoxLayout(self.center_group)
        self.center_layout.setContentsMargins(0, 0, 0, 0)
        self.center_layout.setSpacing(6)

        self.status_badge = StatusBadgeLabel()
        self.btn_toggle_net = QPushButton("Start Net\n(Activate)")
        self.btn_toggle_net.setFixedSize(90, 38)
        self.btn_toggle_net.setStyleSheet(BTN_START_STYLE)
        self.btn_toggle_net.clicked.connect(self.toggle_net_mode)

        self.id_timer_label = DynamicTimerLabel()

        self.center_layout.addWidget(self.btn_toggle_net)
        self.center_layout.addWidget(self.status_badge)
        self.center_layout.addWidget(self.id_timer_label)

        # 3. Right Subgroup (NCS Badge & Net Duration)
        right_subgroup = QHBoxLayout()
        right_subgroup.setContentsMargins(0, 0, 0, 0)
        right_subgroup.setSpacing(6)

        self.ncs_badge = NcsBadgeLabel()
        self.net_duration_label = NetDurationLabel()

        right_subgroup.addWidget(self.ncs_badge)
        right_subgroup.addWidget(self.net_duration_label)

        # Assemble Master Header with Equal Outer Stretches
        header_layout.addLayout(left_subgroup)
        header_layout.addStretch(1)  # Equal left spacer
        header_layout.addWidget(self.center_group)
        header_layout.addStretch(1)  # Equal right spacer
        header_layout.addLayout(right_subgroup)

        main_layout.addLayout(header_layout)

        # Main Splitter
        self.main_splitter = QSplitter(Qt.Orientation.Horizontal)
        self.left_splitter = QSplitter(Qt.Orientation.Vertical)

        # Pane 1: Scripts
        script_group = QGroupBox("Net Script")
        script_layout = QVBoxLayout(script_group)
        script_layout.setContentsMargins(3, 3, 3, 3)

        self.script_tabs = QTabWidget()
        script_layout.addWidget(self.script_tabs)
        self.left_splitter.addWidget(script_group)

        # Pane 2: Traffic Input
        self.traffic_group = QGroupBox("Net Traffic")
        traffic_layout = QVBoxLayout(self.traffic_group)
        traffic_layout.setContentsMargins(3, 3, 3, 3)

        self.t_row1 = QHBoxLayout()
        self.t_row1.setSpacing(3)

        self.lbl_traffic_station = QLabel("Station:")

        self.combo_traffic_call = QComboBox()
        self.combo_traffic_call.setEditable(True)
        self.combo_traffic_call.setLineEdit(UppercaseLineEdit(self.combo_traffic_call))
        self.combo_traffic_call.setPlaceholderText("Callsign")
        self.combo_traffic_call.setFixedWidth(100)

        self.lbl_traffic_type = QLabel("Type:")
        self.combo_traffic_type = QComboBox()
        self.combo_traffic_type.setFixedWidth(150)

        self.t_row1.addWidget(self.lbl_traffic_station)
        self.t_row1.addWidget(self.combo_traffic_call)
        self.t_row1.addSpacing(4)
        self.t_row1.addWidget(self.lbl_traffic_type)
        self.t_row1.addWidget(self.combo_traffic_type)
        self.t_row1.addStretch()

        self.t_row2 = QHBoxLayout()
        self.t_row2.setSpacing(3)

        self.lbl_traffic_time = QLabel("Time:")
        self.input_traffic_time = QLineEdit()
        self.input_traffic_time.setPlaceholderText("HH:MM")
        self.input_traffic_time.setFixedWidth(55)

        self.btn_traffic_now = QPushButton("Now")
        self.btn_traffic_now.clicked.connect(self.set_traffic_time_now)

        self.lbl_traffic_loc = QLabel("Location:")
        self.input_traffic_loc = QLineEdit()
        self.input_traffic_loc.setPlaceholderText("Location / Grid")

        self.t_row2.addWidget(self.lbl_traffic_time)
        self.t_row2.addWidget(self.input_traffic_time)
        self.t_row2.addWidget(self.btn_traffic_now)
        self.t_row2.addSpacing(4)
        self.t_row2.addWidget(self.lbl_traffic_loc)
        self.t_row2.addWidget(self.input_traffic_loc)

        self.input_traffic_note = QTextEdit()
        self.input_traffic_note.setPlaceholderText("Enter description, question details, or weather report notes...")

        t_actions_row = QHBoxLayout()
        self.btn_add_traffic = QPushButton("Log\nTraffic")
        self.btn_add_traffic.setStyleSheet(BTN_GREEN_STYLE)
        self.btn_add_traffic.clicked.connect(self.log_traffic_item)
        t_actions_row.addWidget(self.btn_add_traffic)

        self.btn_clear_traffic_form = QPushButton("Clear\nForm")
        self.btn_clear_traffic_form.clicked.connect(self.clear_traffic_form)
        t_actions_row.addWidget(self.btn_clear_traffic_form)
        
        self.btn_nts_form = QPushButton("NTS\nMessage")
        self.btn_nts_form.setStyleSheet(BTN_CYAN_STYLE)
        self.btn_nts_form.clicked.connect(self.open_nts_dialog)
        t_actions_row.addWidget(self.btn_nts_form)

        self.status_select_widget = QWidget()
        status_select_layout = QHBoxLayout(self.status_select_widget)
        status_select_layout.setContentsMargins(0, 0, 0, 0)
        status_select_layout.setSpacing(3)

        status_vbox = QVBoxLayout()
        status_vbox.setContentsMargins(0, 0, 0, 0)
        status_vbox.setSpacing(1)

        self.lbl_set_status = QLabel("Net Status:")
        self.combo_net_statuses = QComboBox()
        self.combo_net_statuses.setMinimumWidth(130)

        status_vbox.addWidget(self.lbl_set_status)
        status_vbox.addWidget(self.combo_net_statuses)

        self.btn_set_status = QPushButton("Set\nStatus")
        self.btn_set_status.setStyleSheet(BTN_CYAN_STYLE)
        self.btn_set_status.clicked.connect(self.apply_selected_net_status)

        status_select_layout.addLayout(status_vbox)
        status_select_layout.addWidget(self.btn_set_status)

        t_actions_row.addWidget(self.status_select_widget)

        traffic_layout.addLayout(self.t_row1)
        traffic_layout.addLayout(self.t_row2)
        traffic_layout.addWidget(self.input_traffic_note)
        traffic_layout.addLayout(t_actions_row)
        self.left_splitter.addWidget(self.traffic_group)

        # Right Column
        self.right_splitter = QSplitter(Qt.Orientation.Vertical)

        # Pane 3: Roster
        station_group = QGroupBox("Station Management")
        station_layout = QVBoxLayout(station_group)
        station_layout.setContentsMargins(3, 3, 3, 3)

        self.checkin_layout = QHBoxLayout()
        self.checkin_layout.setSpacing(3)

        self.input_callsign = UppercaseLineEdit()
        self.input_callsign.setPlaceholderText("Callsign (Alt+1)")
        self.input_callsign.returnPressed.connect(self.checkin_station)

        self.completer_model = CallsignCompletionModel(self.station_memory, self)
        self.completer = QCompleter(self.completer_model, self)
        self.completer.setCaseSensitivity(Qt.CaseSensitivity.CaseInsensitive)
        self.completer.setFilterMode(Qt.MatchFlag.MatchContains)
        self.completer.setCompletionMode(QCompleter.CompletionMode.PopupCompletion)
        self.input_callsign.setCompleter(self.completer)

        self.lbl_roster_type = QLabel("Type:")
        self.combo_station_type = QComboBox()

        self.chk_traffic = EnterToggleCheckBox("Traffic")
        self.chk_announcement = EnterToggleCheckBox("Announcement")

        self.lbl_roster_loc = QLabel("Info / Grid:")
        self.input_roster_loc = QLineEdit()
        self.input_roster_loc.setPlaceholderText("City / Grid / Info")
        self.input_roster_loc.returnPressed.connect(self.checkin_station)

        btn_checkin = EnterClickButton("Check In")
        btn_checkin.setStyleSheet(BTN_GREEN_STYLE)
        btn_checkin.clicked.connect(self.checkin_station)

        self.checkin_layout.addWidget(self.input_callsign)
        self.checkin_layout.addWidget(self.lbl_roster_type)
        self.checkin_layout.addWidget(self.combo_station_type)
        self.checkin_layout.addWidget(self.chk_traffic)
        self.checkin_layout.addWidget(self.chk_announcement)
        self.checkin_layout.addWidget(self.lbl_roster_loc)
        self.checkin_layout.addWidget(self.input_roster_loc)
        self.checkin_layout.addWidget(btn_checkin)

        self.station_table = QTableWidget(0, 7)
        self.station_table.horizontalHeader().setFixedHeight(22)
        self.station_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.station_table.itemChanged.connect(self.on_table_cell_changed)

        st_actions = QHBoxLayout()
        st_actions.setSpacing(3)

        self.btn_roll_call = QPushButton("Roll\nCall")
        self.btn_roll_call.clicked.connect(self.start_roll_call)

        btn_search_edit_station = QPushButton("Find Call\nEdit Roster")
        btn_search_edit_station.clicked.connect(self.open_combined_station_search_edit)

        btn_toggle_standby = QPushButton("Toggle\nStatus")
        btn_toggle_standby.clicked.connect(self.toggle_station_standby)

        btn_checkout = QPushButton("Check\nOut")
        btn_checkout.setStyleSheet(BTN_ORANGE_STYLE)
        btn_checkout.clicked.connect(self.checkout_station)

        btn_delete_station = QPushButton("Delete\nStation")
        btn_delete_station.setStyleSheet(BTN_DANGER_STYLE)
        btn_delete_station.clicked.connect(self.delete_station)

        btn_clear_roster = QPushButton("Clear\nRoster")
        btn_clear_roster.setStyleSheet(BTN_DANGER_STYLE)
        btn_clear_roster.clicked.connect(self.clear_station_roster)

        counter_frame = QFrame()
        counter_frame.setFrameShape(QFrame.Shape.Panel)
        counter_frame.setFrameShadow(QFrame.Shadow.Sunken)
        counter_frame.setFixedWidth(115)

        counter_layout = QHBoxLayout(counter_frame)
        counter_layout.setContentsMargins(4, 1, 4, 1)

        self.lbl_station_counts = QLabel("Total: 0\nMobile: 0")
        self.lbl_station_counts.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        mono_font = QFont("Consolas", 8, QFont.Weight.Bold)
        mono_font.setStyleHint(QFont.StyleHint.Monospace)
        self.lbl_station_counts.setFont(mono_font)
        self.lbl_station_counts.setStyleSheet("color: #000080;")
        counter_layout.addWidget(self.lbl_station_counts)

        st_actions.addWidget(self.btn_roll_call)
        st_actions.addWidget(btn_search_edit_station)
        st_actions.addWidget(btn_toggle_standby)
        st_actions.addWidget(btn_checkout)
        st_actions.addWidget(btn_delete_station)
        st_actions.addWidget(btn_clear_roster)
        st_actions.addStretch()
        st_actions.addWidget(counter_frame)

        station_layout.addLayout(self.checkin_layout)
        station_layout.addWidget(self.station_table)
        station_layout.addLayout(st_actions)

        self.right_splitter.addWidget(station_group)

        # Pane 4: Activity Log
        self.activity_group = QGroupBox("Traffic and Activity Log")
        activity_layout = QVBoxLayout(self.activity_group)
        activity_layout.setContentsMargins(3, 3, 3, 3)

        self.activity_table = QTableWidget(0, 5)
        self.activity_table.horizontalHeader().setFixedHeight(22)
        self.activity_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.activity_table.itemChanged.connect(self.on_activity_cell_changed)

        act_actions = QHBoxLayout()
        act_actions.setSpacing(3)
        btn_traffic_search = QPushButton("Find Traffic\nEdit Log")
        btn_traffic_search.clicked.connect(self.open_combined_traffic_search_edit)

        btn_traffic_delete = QPushButton("Delete\nEntry")
        btn_traffic_delete.setStyleSheet(BTN_DANGER_STYLE)
        btn_traffic_delete.clicked.connect(self.delete_traffic_entry)

        btn_clear_log = QPushButton("Clear\nLog")
        btn_clear_log.setStyleSheet(BTN_DANGER_STYLE)
        btn_clear_log.clicked.connect(self.clear_activity_log)

        act_actions.addWidget(btn_traffic_search)
        act_actions.addWidget(btn_traffic_delete)
        act_actions.addWidget(btn_clear_log)
        act_actions.addStretch()

        self.lbl_web_status = QLabel("Web Status:")
        self.combo_manual_web_status = QComboBox()
        self.combo_manual_web_status.setEditable(False)
        self.combo_manual_web_status.setMinimumWidth(160)

        self.btn_push_web_status = QPushButton("Send\nStatus")
        self.btn_push_web_status.setStyleSheet(BTN_CYAN_STYLE)
        self.btn_push_web_status.clicked.connect(self.push_custom_web_status)

        act_actions.addWidget(self.lbl_web_status)
        act_actions.addWidget(self.combo_manual_web_status)
        act_actions.addWidget(self.btn_push_web_status)
        act_actions.addSpacing(10)

        self.input_new_ncs = UppercaseLineEdit()
        self.input_new_ncs.setPlaceholderText("New NCS Callsign")
        self.input_new_ncs.setFixedWidth(120)

        btn_transfer_ncs = QPushButton("Transfer\nNet")
        btn_transfer_ncs.setStyleSheet(BTN_ORANGE_STYLE)
        btn_transfer_ncs.clicked.connect(self.transfer_net_control)

        act_actions.addWidget(self.input_new_ncs)
        act_actions.addWidget(btn_transfer_ncs)

        activity_layout.addWidget(self.activity_table)
        activity_layout.addLayout(act_actions)
        self.right_splitter.addWidget(self.activity_group)

        self.left_splitter.setSizes([500, 500])
        self.left_splitter.setStretchFactor(0, 1)
        self.left_splitter.setStretchFactor(1, 1)

        self.right_splitter.setSizes([500, 500])
        self.right_splitter.setStretchFactor(0, 1)
        self.right_splitter.setStretchFactor(1, 1)

        self.main_splitter.addWidget(self.left_splitter)
        self.main_splitter.addWidget(self.right_splitter)
        self.main_splitter.setSizes([380, 860])

        main_layout.addWidget(self.main_splitter, 1)

        # Footer
        footer_layout = QHBoxLayout()
        footer_layout.setSpacing(4)

        self.btn_open_session = QPushButton("Open\nSession")
        self.btn_open_session.clicked.connect(self.open_saved_session_file)

        self.btn_save_session = QPushButton("Save\nSession")
        self.btn_save_session.setStyleSheet(BTN_SAVE_STYLE)
        self.btn_save_session.clicked.connect(self.manual_save_session)

        self.btn_restore = QPushButton("Restore\nAutosave")
        self.btn_restore.setStyleSheet(BTN_RESTORE_STYLE)
        self.btn_restore.clicked.connect(self.restore_session)

        self.footer_legend_frame = QFrame()
        self.footer_legend_frame.setFrameShape(QFrame.Shape.Panel)
        self.footer_legend_frame.setFrameShadow(QFrame.Shadow.Sunken)

        self.footer_legend_layout = QHBoxLayout(self.footer_legend_frame)
        self.footer_legend_layout.setContentsMargins(6, 2, 6, 2)

        self.footer_legend_lbl = QLabel("[F1]\nDisplay Hotkeys")
        self.footer_legend_lbl.setStyleSheet("color: #000080;")
        self.footer_legend_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.footer_legend_layout.addWidget(self.footer_legend_lbl)

        self.repeater_badge = RepeaterBadgeLabel()
        self.utc_clock_badge = ClockBadgeLabel("UTC TIME", is_utc=True, use_24hr=self.use_24hr_clock)
        self.local_clock_badge = ClockBadgeLabel("SYSTEM TIME", is_utc=False, use_24hr=self.use_24hr_clock)

        self.btn_export_xlsx = QPushButton("Export XLSX\nCheck-Ins")
        self.btn_export_xlsx.clicked.connect(self.export_xlsx)

        self.btn_export_csv = QPushButton("Export\nCSV")
        self.btn_export_csv.clicked.connect(self.export_csv)

        self.btn_export_pdf = QPushButton("Export\nPDF")
        self.btn_export_pdf.clicked.connect(self.export_pdf)

        self.btn_reset = QPushButton("Reset\nNet")
        self.btn_reset.setStyleSheet(BTN_DANGER_STYLE)
        self.btn_reset.clicked.connect(self.reset_entire_net)

        self.btn_exit = QPushButton("Exit\nApp")
        self.btn_exit.setStyleSheet(BTN_DANGER_STYLE)
        self.btn_exit.clicked.connect(self.close)

        footer_layout.addWidget(self.btn_open_session)
        footer_layout.addWidget(self.btn_save_session)
        footer_layout.addWidget(self.btn_restore)
        footer_layout.addStretch()
        footer_layout.addWidget(self.footer_legend_frame)
        footer_layout.addStretch()
        footer_layout.addWidget(self.repeater_badge)
        footer_layout.addWidget(self.utc_clock_badge)
        footer_layout.addWidget(self.local_clock_badge)
        footer_layout.addWidget(self.btn_export_xlsx)
        footer_layout.addWidget(self.btn_export_csv)
        footer_layout.addWidget(self.btn_export_pdf)
        footer_layout.addWidget(self.btn_reset)
        footer_layout.addWidget(self.btn_exit)

        main_layout.addLayout(footer_layout)

    def open_nts_dialog(self):
        station = self.combo_traffic_call.currentText().strip().upper() or self.current_ncs_callsign
        dialog = NtsRadiogramDialog(default_station="", parent=self)

        if dialog.exec() == QDialog.DialogCode.Accepted:
            data = dialog.get_radiogram_data()
            summary_note = f"ARRL RADIOGRAM #{data['num']} ({data['precedence']}) - TO: {data['to_name']} - CK: {data['check']}"

            # Log to Activity Table
            self.add_activity_log(
                call=data["origin"],
                traffic_type="Priority" if data["precedence"] in ("P", "E") else "General Note",
                location=data["place"],
                note=summary_note
            )

            # Auto-export individual Radiogram PDF
            default_pdf = f"Radiogram_{data['num']}_{data['origin']}.pdf"
            path, _ = QFileDialog.getSaveFileName(self, "Save Radiogram PDF", default_pdf, "PDF Files (*.pdf)")
            if path:
                dialog.export_radiogram_pdf(path, data)
                QMessageBox.information(self, "PDF Exported", f"Radiogram saved to:\n{path}")

    def on_table_cell_changed(self, item):
        row = item.row()
        col = item.column()
        if row < 0 or row >= len(self.stations):
            return

        self.station_table.blockSignals(True)
        try:
            new_text = item.text().strip()
            st = self.stations[row]

            if self.current_mode == "SKYWARN":
                if col == 0:
                    st["callsign"] = new_text.upper()
                    item.setText(st["callsign"])
                    self.register_station_callsign(st["callsign"])
                    self.refresh_callsign_combo()
                elif col == 1:
                    st["type"] = new_text
                elif col == 2:
                    st["detail"] = new_text
                elif col == 3:
                    st["status"] = new_text
                elif col == 4:
                    st["time_in"] = new_text
                elif col == 5:
                    st["time_out"] = new_text
            else:
                if col == 0:
                    st["callsign"] = new_text.upper()
                    item.setText(st["callsign"])
                    self.register_station_callsign(st["callsign"], st.get("info", ""))
                    self.refresh_callsign_combo()
                elif col == 1:
                    st["type"] = new_text
                elif col == 2:
                    st["detail"] = new_text
                elif col == 3:
                    st["info"] = new_text
                elif col == 4:
                    st["status"] = new_text
                elif col == 5:
                    st["time_in"] = new_text
                elif col == 6:
                    st["time_out"] = new_text
        finally:
            self.station_table.blockSignals(False)

    def open_combined_station_search_edit(self):
        prof = self.profiles.get(self.current_profile_name, {})
        st_types = prof.get("station_types", ["Base", "Mobile", "NCS"])

        dialog = CombinedStationSearchEditDialog(
            stations=self.stations,
            station_types=st_types,
            is_skywarn=(self.current_mode == "SKYWARN"),
            parent=self
        )
        dialog.exec()
        self.refresh_station_table()
        self.refresh_callsign_combo()

    def focus_web_status_dropdown(self):
        if self.combo_manual_web_status.isVisible():
            self.combo_manual_web_status.setFocus()

    def focus_new_ncs_input(self):
        self.input_new_ncs.setFocus()
        self.input_new_ncs.selectAll()

    def clear_traffic_form(self):
        self.combo_traffic_call.setCurrentText("")
        self.input_traffic_time.clear()
        self.input_traffic_loc.clear()
        self.input_traffic_note.clear()
        self.combo_traffic_call.setFocus()

    def push_custom_web_status(self):
        if not self.combo_manual_web_status.isVisible():
            return

        custom_status = self.combo_manual_web_status.currentText().strip()
        if not custom_status:
            QMessageBox.warning(self, "Web Status Error", "Please select a web status to send.")
            return

        prof = self.profiles.get(self.current_profile_name, {})
        api_cfg = prof.get("api_sync", {})

        if not api_cfg or not api_cfg.get("enabled", False):
            QMessageBox.warning(self, "API Sync Disabled", "Online API sync is not enabled for the current profile.")
            return

        dispatch_api_status_update(api_cfg, custom_status, app_ref=self)
        self.web_status_badge.set_status(custom_status)
        
        # Use active operator callsign
        active_call = self.get_active_logging_callsign()
        self.add_activity_log(active_call, "API UPDATE", "-", f"=== MANUAL WEB STATUS PUSH: {custom_status} ===")

    def auto_add_ncs_to_roster(self, callsign):
        call = callsign.strip().upper()
        if not call or call == "UNASSIGNED":
            return

        prof = self.profiles.get(self.current_profile_name, {})
        selectable_types = prof.get("station_types", [])

        matched_ncs_type = None
        for t in selectable_types:
            t_upper = t.strip().upper()
            if t_upper in ("NCS", "NET CONTROL", "NCO"):
                matched_ncs_type = t
                break

        if not matched_ncs_type:
            matched_ncs_type = selectable_types[0] if selectable_types else "NCS"

        self.register_station_callsign(call, "Net Control Station")
        now_str = datetime.now().strftime("%H:%M")

        for st in self.stations:
            if st["callsign"] == call:
                st["type"] = matched_ncs_type
                st["status"] = "Active"
                st["time_in"] = now_str
                st["time_out"] = ""
                self.refresh_station_table()
                self.refresh_callsign_combo()
                return

        col_detail = "-"
        col_info = "Net Control Station"

        station_data = {
            "callsign": call,
            "type": matched_ncs_type,
            "detail": col_detail,
            "info": col_info,
            "status": "Active",
            "time_in": now_str,
            "time_out": ""
        }
        self.stations.append(station_data)
        self.refresh_station_table()
        self.refresh_callsign_combo()

    def auto_add_scribe_to_roster(self, callsign, station_type="Alternate NCS"):
        call = callsign.strip().upper()
        if not call:
            return

        now_str = datetime.now().strftime("%H:%M")
        self.register_station_callsign(call, "Net Scribe")

        for st in self.stations:
            if st["callsign"] == call:
                st["type"] = station_type
                st["status"] = "Active"
                st["time_in"] = now_str
                st["time_out"] = ""
                self.refresh_station_table()
                self.refresh_callsign_combo()
                return

        col_detail = "-"
        col_info = "Net Scribe / Assistant"

        station_data = {
            "callsign": call,
            "type": station_type,
            "detail": col_detail,
            "info": col_info,
            "status": "Active",
            "time_in": now_str,
            "time_out": ""
        }
        self.stations.append(station_data)
        self.refresh_station_table()
        self.refresh_callsign_combo()

    def transfer_net_control(self):
        new_ncs = self.input_new_ncs.text().strip().upper()
        if not new_ncs:
            QMessageBox.warning(self, "Net Transfer Error", "Please enter a valid callsign for the new Net Control Station.")
            return

        old_ncs = self.current_ncs_callsign
        if new_ncs == old_ncs:
            QMessageBox.information(self, "Transfer Net", f"{new_ncs} is already assigned as Net Control.")
            return

        confirm = QMessageBox.question(
            self, "Confirm Net Transfer",
            f"Transfer Net Control responsibility from {old_ncs} to {new_ncs}?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.Yes
        )

        if confirm == QMessageBox.StandardButton.Yes:
            self.current_ncs_callsign = new_ncs
            self.ncs_badge.set_ncs(new_ncs, self.current_scribe_callsign)
            self.input_new_ncs.clear()

            prof = self.profiles.get(self.current_profile_name, {})
            if "api_sync" in prof:
                prof["api_sync"]["qrz"] = new_ncs

            self.auto_add_ncs_to_roster(new_ncs)

            self.add_activity_log(
                call="NET CONTROL",
                traffic_type="General Note",
                location="-",
                note=f"=== NET CONTROL TRANSFERRED FROM {old_ncs} TO {new_ncs} ==="
            )

    def register_station_callsign(self, callsign, name=""):
        call = callsign.strip().upper()
        if not call:
            return

        existing = next((item for item in self.station_memory if item.get("callsign", "").upper() == call), None)
        if existing:
            if name and not existing.get("name"):
                existing["name"] = name
        else:
            self.station_memory.append({"callsign": call, "name": name})

        save_station_memory(self.station_memory)
        self.completer_model.set_items(self.station_memory)

    def add_rollcall_checkin(self, callsign, detail, info_name):
        self.register_station_callsign(callsign, info_name)
        now_str = datetime.now().strftime("%H:%M")
        default_type = self.combo_station_type.itemText(0) if self.combo_station_type.count() > 0 else "Base"

        for st in self.stations:
            if st["callsign"] == callsign:
                st["detail"] = detail
                st["info"] = info_name or st.get("info", "-")
                st["status"] = "Active"
                st["time_in"] = now_str
                st["time_out"] = ""
                self.refresh_station_table()
                self.refresh_callsign_combo()
                self.station_table.scrollToBottom()
                return

        st_data = {
            "callsign": callsign,
            "type": default_type,
            "detail": detail,
            "info": info_name or "-",
            "status": "Active",
            "time_in": now_str,
            "time_out": ""
        }
        self.stations.append(st_data)
        self.refresh_station_table()
        self.refresh_callsign_combo()
        self.station_table.scrollToBottom()

    def start_roll_call(self):
        if not self.has_roll_call:
            return

        # Check if we have an active roll call in memory
        if getattr(self, "rollcall_roster", []):
            reply = QMessageBox.question(
                self, "Resume Roll Call",
                "A roll call is already in progress. Do you want to resume it?\n\nSelect 'No' to load a new roster file.",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                QMessageBox.StandardButton.Yes
            )
            if reply == QMessageBox.StandardButton.Yes:
                dialog = RollCallDialog(self.rollcall_roster, start_index=self.rollcall_index, parent=self)
                dialog.exec()
                self.rollcall_index = dialog.index
                return

        # If no resume, prompt for a new file
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Select Roll Call File", "", "Excel Files (*.xlsx *.xls)"
        )
        if not file_path:
            return

        try:
            roster = parse_roll_call_file(file_path)
            if not roster:
                QMessageBox.warning(self, "Roll Call Error", "No valid callsigns were found in the selected file.")
                return

            self.rollcall_roster = roster
            self.rollcall_index = 0
            dialog = RollCallDialog(self.rollcall_roster, start_index=self.rollcall_index, parent=self)
            dialog.exec()
            self.rollcall_index = dialog.index
        except Exception as e:
            QMessageBox.critical(self, "Roll Call Error", f"Failed to parse roll call file:\n{str(e)}")

    def update_api_ui_elements(self, is_api_enabled):
        self.web_status_badge.setVisible(is_api_enabled)
        self.lbl_web_status.setVisible(is_api_enabled)
        self.combo_manual_web_status.setVisible(is_api_enabled)
        self.btn_push_web_status.setVisible(is_api_enabled)

        while self.center_layout.count() > 0:
            item = self.center_layout.takeAt(0)
            if item.widget():
                item.widget().setParent(None)

        if is_api_enabled:
            self.center_layout.addWidget(self.btn_toggle_net)
            self.center_layout.addWidget(self.status_badge)
            self.center_layout.addWidget(self.id_timer_label)
        else:
            self.center_layout.addWidget(self.status_badge)
            self.center_layout.addWidget(self.btn_toggle_net)
            self.center_layout.addWidget(self.id_timer_label)

    def rebuild_ui_for_mode(self, mode_name, traffic_types, net_statuses, has_roll_call=False, use_custom_statuses=False):
        self.current_mode = mode_name
        self.net_statuses = net_statuses
        self.has_roll_call = has_roll_call
        self.use_custom_statuses = use_custom_statuses

        self.btn_roll_call.setVisible(has_roll_call)
        self.shortcut_roll_call.setEnabled(has_roll_call)

        self.combo_traffic_type.clear()
        self.combo_traffic_type.addItems(traffic_types)

        self.station_table.clear()
        self.activity_table.clear()

        show_status_controls = use_custom_statuses or (mode_name == "SKYWARN")
        self.status_select_widget.setVisible(show_status_controls)

        if show_status_controls:
            self.combo_net_statuses.clear()
            self.combo_net_statuses.addItems(list(net_statuses.keys()))

        prof = self.profiles.get(self.current_profile_name, {})
        web_opts = prof.get("web_status_options", ["ACTIVE", "STANDBY", "CLOSED", "-Normal Operations-"])
        self.combo_manual_web_status.clear()
        self.combo_manual_web_status.addItems(web_opts)
        
        prof = self.profiles.get(self.current_profile_name, {})
        has_nts = prof.get("has_nts", False) and (self.current_mode != "SKYWARN")
        self.btn_nts_form.setVisible(has_nts)

        rep_cfg = prof.get("repeater", {})
        self.repeater_badge.set_repeater_info(rep_cfg, is_using_backup=False)
        

        if mode_name == "SKYWARN":
            self.traffic_group.setTitle("Spotter Reports and Net Traffic")
            self.activity_group.setTitle("SKYWARN Activity Log")

            self.lbl_traffic_time.setVisible(True)
            self.input_traffic_time.setVisible(True)
            self.btn_traffic_now.setVisible(True)
            self.lbl_traffic_loc.setVisible(True)
            self.input_traffic_loc.setVisible(True)
            self.lbl_traffic_type.setText("Condition:")

            self.chk_traffic.setVisible(False)
            self.chk_announcement.setVisible(False)
            self.lbl_roster_loc.setText("Location:")

            self.station_table.setColumnCount(6)
            self.station_table.setHorizontalHeaderLabels(["Callsign", "Station Type", "Location", "Status", "In", "Out"])

            self.station_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Fixed)
            self.station_table.setColumnWidth(0, 60)

            self.station_table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Interactive)
            self.station_table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)

            self.station_table.horizontalHeader().setSectionResizeMode(3, QHeaderView.ResizeMode.Fixed)
            self.station_table.setColumnWidth(3, 90)

            self.station_table.horizontalHeader().setSectionResizeMode(4, QHeaderView.ResizeMode.Fixed)
            self.station_table.setColumnWidth(4, 50)

            self.station_table.horizontalHeader().setSectionResizeMode(5, QHeaderView.ResizeMode.Fixed)
            self.station_table.setColumnWidth(5, 50)

            self.activity_table.setColumnCount(5)
            self.activity_table.setHorizontalHeaderLabels(["Time", "Condition", "Location", "Source", "Narrative"])
            self.activity_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
            self.activity_table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
            self.activity_table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
            self.activity_table.horizontalHeader().setSectionResizeMode(3, QHeaderView.ResizeMode.ResizeToContents)
            self.activity_table.horizontalHeader().setSectionResizeMode(4, QHeaderView.ResizeMode.Stretch)

        else:
            self.traffic_group.setTitle(f"Net Traffic")
            self.activity_group.setTitle(f"Traffic and Activity Log")

            self.lbl_traffic_time.setVisible(False)
            self.input_traffic_time.setVisible(False)
            self.btn_traffic_now.setVisible(False)
            self.lbl_traffic_loc.setVisible(False)
            self.input_traffic_loc.setVisible(False)
            self.lbl_traffic_type.setText("Type:")

            self.chk_traffic.setVisible(True)
            self.chk_announcement.setVisible(True)
            self.lbl_roster_loc.setText("Info / Grid:")

            self.station_table.setColumnCount(7)
            self.station_table.setHorizontalHeaderLabels(["Callsign", "Station Type", "Additional", "Info / Grid", "Status", "In", "Out"])

            self.station_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Fixed)
            self.station_table.setColumnWidth(0, 60)

            self.station_table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Interactive)

            self.station_table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.Fixed)
            self.station_table.setColumnWidth(2, 110)

            self.station_table.horizontalHeader().setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)

            self.station_table.horizontalHeader().setSectionResizeMode(4, QHeaderView.ResizeMode.Fixed)
            self.station_table.setColumnWidth(4, 90)

            self.station_table.horizontalHeader().setSectionResizeMode(5, QHeaderView.ResizeMode.Fixed)
            self.station_table.setColumnWidth(5, 50)

            self.station_table.horizontalHeader().setSectionResizeMode(6, QHeaderView.ResizeMode.Fixed)
            self.station_table.setColumnWidth(6, 50)

            self.activity_table.setColumnCount(4)
            self.activity_table.setHorizontalHeaderLabels(["Time", "Callsign", "Traffic Tag", "Details"])
            self.activity_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
            self.activity_table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
            self.activity_table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
            self.activity_table.horizontalHeader().setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)

        self.refresh_station_table()

    def set_traffic_time_now(self):
        if self.current_mode == "SKYWARN":
            self.input_traffic_time.setText(datetime.now().strftime("%H:%M"))

    def apply_selected_net_status(self):
        if not self.status_select_widget.isVisible():
            return

        st_name = self.combo_net_statuses.currentText()
        if not st_name:
            return

        val = self.net_statuses.get(st_name, {})
        t_hex, b_hex = parse_color_map(val)

        self.net_state_name = st_name
        self.status_badge.set_state(st_name, t_hex, b_hex)
        
        # Use active operator callsign
        active_call = self.get_active_logging_callsign()
        self.add_activity_log(active_call, "STATUS CHANGE", "-", f"=== NET STATUS SET TO {st_name} ===")

        prof = self.profiles.get(self.current_profile_name, {})
        api_cfg = prof.get("api_sync", {})
        dispatch_api_status_update(api_cfg, st_name, app_ref=self)

    def check_profile_access(self, profile_name):
        prof = self.profiles.get(profile_name, {})
        auth_cfg = prof.get("auth", {})

        if not auth_cfg.get("enabled", False):
            return True

        users = auth_cfg.get("users", [])
        if not users:
            return True

        while True:
            auth_dialog = ProfileAuthDialog(profile_name, parent=self)
            if auth_dialog.exec() == QDialog.DialogCode.Accepted:
                call, pin = auth_dialog.get_credentials()
                for u in users:
                    u_call = u.get("callsign", "").strip().upper()
                    u_pass = u.get("passcode", "").strip()
                    if u_call == call and u_pass == pin:
                        return True
                QMessageBox.critical(self, "Access Denied", "Invalid Callsign or Passcode / PIN entered.")
            else:
                return False

    def prompt_profile_on_startup(self):
        self.profiles = load_profiles_from_disk()

        if not self.profiles:
            QMessageBox.information(self, "No Profiles", "No profiles found. Please create a profile to continue.")
            self.open_profile_editor_dialog()
            self.profiles = load_profiles_from_disk()

            if not self.profiles:
                sys.exit(0)

        self.select_profile_and_login()

    def get_active_logging_callsign(self):
        if getattr(self, "current_scribe_callsign", ""):
            return self.current_scribe_callsign
        return self.current_ncs_callsign if self.current_ncs_callsign else "NET CONTROL"

    def prompt_ncs_login(self):
        prof = self.profiles.get(self.current_profile_name, {})
        st_types = prof.get("station_types", ["Base", "Mobile", "NCS", "Alternate NCS"])

        while True:
            # Pass default_ncs="" so input_ncs shows placeholder ghost text instead of hardcoded text
            dialog = NcsScribeLoginDialog(default_ncs="", station_types=st_types, parent=self)
            if dialog.exec() == QDialog.DialogCode.Accepted:
                ncs_call, scribe_call, scribe_type, role = dialog.get_logins()
                self.current_ncs_callsign = ncs_call
                self.current_scribe_callsign = scribe_call
                self.current_user_role = role
                self.ncs_badge.set_ncs(ncs_call, scribe_call)

                if "api_sync" in prof:
                    prof["api_sync"]["qrz"] = ncs_call

                self.auto_add_ncs_to_roster(ncs_call)
                if scribe_call:
                    self.auto_add_scribe_to_roster(scribe_call, station_type=scribe_type)
                return True
            else:
                reply = QMessageBox.question(
                    self,
                    "NCS Required",
                    "A Net Control Station (NCS) is required to operate a net.\n\nDo you wish to add one now?",
                    QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                    QMessageBox.StandardButton.Yes
                )
                if reply == QMessageBox.StandardButton.Yes:
                    continue
                else:
                    return False

    def apply_profile(self, profile_name):
        if profile_name not in self.profiles:
            return

        self.stations.clear()
        self.activities.clear()
        self.activity_table.setRowCount(0)
        self.rollcall_roster = []
        self.rollcall_index = 0

        self.id_timer_label.reset_timer(is_net_active=False)
        self.net_duration_label.reset_duration()

        self.is_net_active = False
        self.btn_toggle_net.setText("Start Net\n(Activate)")
        self.btn_toggle_net.setStyleSheet(BTN_START_STYLE)

        self.current_profile_name = profile_name
        self.profile_badge.set_profile(profile_name)

        prof = self.profiles[profile_name]
        mode = prof.get("mode", "Regular")
        has_rc = prof.get("has_roll_call", False)
        use_cust_st = prof.get("use_custom_statuses", mode == "SKYWARN")
        self.use_24hr_clock = prof.get("use_24hr_clock", False)
        tr_types = prof.get("traffic_types", ["Question", "Announcement", "Priority", "Emergency", "General Note"])
        st_map = prof.get("net_statuses", {})

        self.utc_clock_badge.set_24hr_mode(self.use_24hr_clock)
        self.local_clock_badge.set_24hr_mode(self.use_24hr_clock)

        api_cfg = prof.get("api_sync", {})
        is_api_enabled = api_cfg.get("enabled", False)
        self.update_api_ui_elements(is_api_enabled)

        self.rebuild_ui_for_mode(mode, tr_types, st_map, has_roll_call=has_rc, use_custom_statuses=use_cust_st)

        self.combo_station_type.clear()
        self.combo_station_type.addItems(prof.get("station_types", []))

        self.script_tabs.clear()
        for sc in prof.get("scripts", []):
            self.script_tabs.addTab(
                self.create_script_tab(sc.get("text", "")),
                sc.get("title", "Script")
            )

        self.net_state_name = "INACTIVE"
        self.status_badge.set_state("INACTIVE", "#FFFF00", "#000000")

        self.refresh_callsign_combo()
        self.main_splitter.setSizes([380, 860])

        self.poll_remote_web_status()
        
    def select_profile_and_login(self):
        while True:
            self.profiles = load_profiles_from_disk()
            dialog = ProfileSelectDialog(self.profiles, self.current_profile_name, self)
            
            if dialog.exec() == QDialog.DialogCode.Accepted and dialog.selected_profile:
                target_prof = dialog.selected_profile
                if self.check_profile_access(target_prof):
                    self.apply_profile(target_prof)
                    if self.prompt_ncs_login():
                        break
            else:
                break

    def open_profile_select_dialog(self):
        self.select_profile_and_login()

    def open_profile_editor_dialog(self):
        self.profiles = load_profiles_from_disk()
        dialog = ProfileEditorDialog(self.profiles, self.current_profile_name, self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.profiles = load_profiles_from_disk()
            if self.current_profile_name in self.profiles:
                self.apply_profile(self.current_profile_name)
            elif self.profiles:
                self.apply_profile(list(self.profiles.keys())[0])

    def focus_callsign_input(self):
        self.input_callsign.setFocus()
        self.input_callsign.selectAll()

    def focus_traffic_station_input(self):
        self.combo_traffic_call.setFocus()

    def next_script_tab(self):
        count = self.script_tabs.count()
        if count > 0:
            curr = self.script_tabs.currentIndex()
            self.script_tabs.setCurrentIndex((curr + 1) % count)

    def prev_script_tab(self):
        count = self.script_tabs.count()
        if count > 0:
            curr = self.script_tabs.currentIndex()
            self.script_tabs.setCurrentIndex((curr - 1 + count) % count)

    def delete_station(self):
        focused = QApplication.focusWidget()
        if isinstance(focused, (QLineEdit, QTextEdit)):
            return

        selected = self.station_table.currentRow()
        if selected < 0 or selected >= len(self.stations):
            QMessageBox.information(self, "Delete Station", "Please select a station from the roster to delete.")
            return

        st = self.stations[selected]
        dialog = RedDangerDialog("Confirm Delete Station", f"Delete station '{st['callsign']}' from roster?", require_type_delete=True, parent=self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.stations.pop(selected)
            self.refresh_station_table()
            self.refresh_callsign_combo()

    def clear_station_roster(self):
        if not self.stations:
            return

        dialog = RedDangerDialog("Confirm Clear Roster", "Clear all checked-in stations from the roster?", require_type_delete=False, parent=self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.stations.clear()
            self.refresh_station_table()
            self.refresh_callsign_combo()

    def open_combined_traffic_search_edit(self):
        prof = self.profiles.get(self.current_profile_name, {})
        tr_types = prof.get("traffic_types", [])

        dialog = CombinedTrafficSearchEditDialog(
            activities=self.activities,
            traffic_types=tr_types,
            is_skywarn=(self.current_mode == "SKYWARN"),
            parent=self
        )
        dialog.exec()
        self.refresh_activity_table_display()

    def refresh_activity_table_display(self):
        self.activity_table.blockSignals(True)
        try:
            self.activity_table.setRowCount(0)
            for act in self.activities:
                row = self.activity_table.rowCount()
                self.activity_table.insertRow(row)

                if self.current_mode == "SKYWARN":
                    self.activity_table.setItem(row, 0, QTableWidgetItem(act.get("timestamp", "")))
                    self.activity_table.setItem(row, 1, QTableWidgetItem(act.get("condition", act.get("type", ""))))
                    self.activity_table.setItem(row, 2, QTableWidgetItem(act.get("location", "-")))
                    self.activity_table.setItem(row, 3, QTableWidgetItem(act.get("source", act.get("callsign", ""))))
                    self.activity_table.setItem(row, 4, QTableWidgetItem(act.get("narrative", act.get("note", ""))))
                else:
                    self.activity_table.setItem(row, 0, QTableWidgetItem(act.get("timestamp", "")))
                    self.activity_table.setItem(row, 1, QTableWidgetItem(act.get("callsign", "")))

                    type_item = QTableWidgetItem(act.get("type", ""))
                    if act.get("type") == "Emergency":
                        type_item.setForeground(QColor("#FF0000"))
                    elif act.get("type") == "Priority":
                        type_item.setForeground(QColor("#804000"))

                    self.activity_table.setItem(row, 2, type_item)
                    self.activity_table.setItem(row, 3, QTableWidgetItem(act.get("note", "")))
        finally:
            self.activity_table.blockSignals(False)

    def delete_traffic_entry(self):
        selected = self.activity_table.currentRow()
        if selected < 0 or selected >= len(self.activities):
            QMessageBox.information(self, "Delete Traffic", "Please select a traffic entry to delete.")
            return

        call = self.activities[selected].get("callsign", self.activities[selected].get("source", ""))
        dialog = RedDangerDialog("Confirm Delete Traffic Entry", f"Delete selected traffic entry for {call}?", require_type_delete=True, parent=self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.activities.pop(selected)
            self.activity_table.removeRow(selected)

    def clear_activity_log(self):
        if not self.activities:
            return

        dialog = RedDangerDialog("Confirm Clear Log", "Clear the entire activity log?", require_type_delete=False, parent=self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.activities.clear()
            self.activity_table.setRowCount(0)

    def reset_entire_net(self):
        dialog = RedDangerDialog("Confirm Net Reset", "Reset net? This clears all stations, logs, and timers.", require_type_delete=False, parent=self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.stations.clear()
            self.activities.clear()
            self.rollcall_roster = []
            self.rollcall_index = 0
            self.refresh_station_table()
            self.refresh_callsign_combo()
            self.activity_table.setRowCount(0)

            self.id_timer_label.reset_timer(is_net_active=False)
            self.net_duration_label.reset_duration()

            if self.is_net_active:
                self.toggle_net_mode()
            else:
                self.net_state_name = "INACTIVE"
                self.status_badge.set_state("INACTIVE", "#FFFF00", "#000000")

            self.current_ncs_callsign = "UNASSIGNED"
            self.current_scribe_callsign = ""
            self.ncs_badge.set_ncs("UNASSIGNED")

            if not self.prompt_ncs_login():
                self.select_profile_and_login()

    def autosave_session(self):
        data = {
            "is_net_active": self.is_net_active,
            "net_state_name": self.net_state_name,
            "current_ncs_callsign": self.current_ncs_callsign,
            "current_scribe_callsign": self.current_scribe_callsign,
            "stations": self.stations,
            "activities": self.activities,
            "time_remaining": self.id_timer_label.time_remaining,
            "elapsed_seconds": self.net_duration_label.elapsed_seconds,
            "start_time_str": self.net_duration_label.start_time_str,
            "end_time_str": self.net_duration_label.end_time_str,
            "current_profile": self.current_profile_name,
            "base_font_size": self.base_font_size,
            "is_using_backup": self.repeater_badge.is_using_backup,
            "rollcall_roster": getattr(self, "rollcall_roster", []),
            "rollcall_index": getattr(self, "rollcall_index", 0)
        }
        try:
            with open(self.AUTOSAVE_FILE, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=2)
        except Exception as e:
            print(f"Autosave Error: {e}")

    def manual_save_session(self):
        path, _ = QFileDialog.getSaveFileName(self, "Save Net Session State", "net_session.json", "JSON Session Files (*.json)")
        if not path:
            return

        data = {
            "is_net_active": self.is_net_active,
            "net_state_name": self.net_state_name,
            "current_ncs_callsign": self.current_ncs_callsign,
            "current_scribe_callsign": self.current_scribe_callsign,
            "stations": self.stations,
            "activities": self.activities,
            "time_remaining": self.id_timer_label.time_remaining,
            "elapsed_seconds": self.net_duration_label.elapsed_seconds,
            "start_time_str": self.net_duration_label.start_time_str,
            "end_time_str": self.net_duration_label.end_time_str,
            "current_profile": self.current_profile_name,
            "base_font_size": self.base_font_size,
            "is_using_backup": self.repeater_badge.is_using_backup
        }
        try:
            with open(path, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=2)
            QMessageBox.information(self, "Save Session", f"Session successfully saved to:\n{path}")
        except Exception as e:
            QMessageBox.critical(self, "Save Error", f"Failed to save session state:\n{str(e)}")

    def open_saved_session_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Open Saved Net Session State", "", "JSON Session Files (*.json)")
        if not path or not os.path.exists(path):
            return

        self.load_session_from_dict(path)

    def restore_session(self):
        if not os.path.exists(self.AUTOSAVE_FILE):
            QMessageBox.warning(self, "Restore Session", "No autosave session file found.")
            return

        self.load_session_from_dict(self.AUTOSAVE_FILE)

    def load_session_from_dict(self, file_path):
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                data = json.load(f)

            self.is_net_active = data.get("is_net_active", False)
            self.net_state_name = data.get("net_state_name", "INACTIVE")
            self.current_ncs_callsign = data.get("current_ncs_callsign", "UNASSIGNED")
            self.current_scribe_callsign = data.get("current_scribe_callsign", "")
            self.ncs_badge.set_ncs(self.current_ncs_callsign, self.current_scribe_callsign)

            self.stations = data.get("stations", [])
            self.activities = data.get("activities", [])
            self.rollcall_roster = data.get("rollcall_roster", [])
            self.rollcall_index = data.get("rollcall_index", 0)

            saved_font_size = data.get("base_font_size", 10)
            self.base_font_size = saved_font_size
            self.apply_font_size()

            saved_prof = data.get("current_profile", "")
            if saved_prof and saved_prof in self.profiles:
                self.current_profile_name = saved_prof
                self.profile_badge.set_profile(saved_prof)

                prof = self.profiles[saved_prof]
                mode = prof.get("mode", "Regular")
                has_rc = prof.get("has_roll_call", False)
                use_cust_st = prof.get("use_custom_statuses", mode == "SKYWARN")
                self.use_24hr_clock = prof.get("use_24hr_clock", False)
                tr_types = prof.get("traffic_types", ["Question", "Announcement", "Priority", "Emergency", "General Note"])
                st_map = prof.get("net_statuses", {})

                self.utc_clock_badge.set_24hr_mode(self.use_24hr_clock)
                self.local_clock_badge.set_24hr_mode(self.use_24hr_clock)

                api_cfg = prof.get("api_sync", {})
                is_api_enabled = api_cfg.get("enabled", False)
                self.update_api_ui_elements(is_api_enabled)

                self.rebuild_ui_for_mode(mode, tr_types, st_map, has_roll_call=has_rc, use_custom_statuses=use_cust_st)

                # Restore repeater state (primary vs. backup)
                is_using_backup = data.get("is_using_backup", False)
                rep_cfg = prof.get("repeater", {})
                self.repeater_badge.set_repeater_info(rep_cfg, is_using_backup=is_using_backup)

                self.combo_station_type.clear()
                self.combo_station_type.addItems(prof.get("station_types", []))

                self.script_tabs.clear()
                for sc in prof.get("scripts", []):
                    self.script_tabs.addTab(
                        self.create_script_tab(sc.get("text", "")),
                        sc.get("title", "Script")
                    )

            time_rem = data.get("time_remaining", 600)
            self.id_timer_label.time_remaining = time_rem
            self.id_timer_label.update_display_text()

            self.net_duration_label.elapsed_seconds = data.get("elapsed_seconds", 0)
            self.net_duration_label.start_time_str = data.get("start_time_str", "--:--")
            self.net_duration_label.end_time_str = data.get("end_time_str", "--:--")
            self.net_duration_label.update_display()

            val = self.net_statuses.get(self.net_state_name, {})
            t_hex, b_hex = parse_color_map(val)
            self.status_badge.set_state(self.net_state_name, t_hex, b_hex)

            if self.is_net_active:
                self.btn_toggle_net.setText("Close Net")
                self.btn_toggle_net.setStyleSheet(BTN_CLOSE_STYLE)
                if not self.net_duration_label.duration_timer.isActive():
                    self.net_duration_label.duration_timer.start()
                self.id_timer_label.countdown_timer.start()
            else:
                self.btn_toggle_net.setText("Start Net\n(Activate)")
                self.btn_toggle_net.setStyleSheet(BTN_START_STYLE)
                self.net_duration_label.stop_duration()
                self.id_timer_label.countdown_timer.stop()

            self.refresh_station_table()
            self.refresh_callsign_combo()
            self.refresh_activity_table_display()

            QMessageBox.information(self, "Session Loaded", "Session state successfully restored!")
        except Exception as e:
            QMessageBox.critical(self, "Load Error", f"Failed to load session state:\n{str(e)}")

    def create_script_tab(self, default_text):
        editor = QTextEdit()
        formatted_text = str(default_text).replace("\\n", "\n")
        editor.setPlainText(formatted_text)
        editor.setReadOnly(True)
        return editor

    def toggle_net_mode(self):
        self.is_net_active = not self.is_net_active
        prof = self.profiles.get(self.current_profile_name, {})
        api_cfg = prof.get("api_sync", {})
        
        # Determine active operator callsign (e.g. W9JBA)
        active_call = self.get_active_logging_callsign()

        if self.is_net_active:
            if self.status_select_widget.isVisible() and self.combo_net_statuses.count() > 0:
                selected_st = self.combo_net_statuses.currentText()
                val = self.net_statuses.get(selected_st, {})
                t_hex, b_hex = parse_color_map(val)
                self.net_state_name = selected_st
                self.status_badge.set_state(selected_st, t_hex, b_hex)
            else:
                self.net_state_name = "ACTIVE"
                val = self.net_statuses.get("ACTIVE", {})
                t_hex, b_hex = parse_color_map(val)
                self.status_badge.set_state("ACTIVE", t_hex, b_hex)

            self.btn_toggle_net.setText("Close Net")
            self.btn_toggle_net.setStyleSheet(BTN_CLOSE_STYLE)
            
            # Use active_call instead of hardcoded "NET CONTROL"
            self.add_activity_log(active_call, "General Note", "-", f"=== NET SESSION STARTED ({self.net_state_name}) ===")

            self.id_timer_label.reset_timer(is_net_active=True)
            self.net_duration_label.reset_duration()
            self.net_duration_label.start_duration()

            dispatch_api_status_update(api_cfg, self.net_state_name, app_ref=self)
        else:
            self.net_state_name = "CLOSED"
            val = self.net_statuses.get("CLOSED", {})
            t_hex, b_hex = parse_color_map(val)

            self.status_badge.set_state("CLOSED", t_hex, b_hex)
            self.btn_toggle_net.setText("Start Net\n(Activate)")
            self.btn_toggle_net.setStyleSheet(BTN_START_STYLE)
            
            # Use active_call instead of hardcoded "NET CONTROL"
            self.add_activity_log(active_call, "General Note", "-", "=== NET SESSION PLACED ON STANDBY / CLOSED ===")

            self.id_timer_label.countdown_timer.stop()
            self.net_duration_label.stop_duration()

            dispatch_api_status_update(api_cfg, False, app_ref=self)

    def format_additional_attributes(self, has_traffic, has_announcement):
        items = []
        if has_traffic:
            items.append("Traffic")
        if has_announcement:
            items.append("Announcement")
        return ", ".join(items) if items else "-"

    def checkin_station(self):
        raw_call = self.input_callsign.text().strip().upper()
        call = raw_call.split("—")[0].strip()
        st_type = self.combo_station_type.currentText()
        now_str = datetime.now().strftime("%H:%M")
        info_val = self.input_roster_loc.text().strip() or "-"

        if not call:
            QMessageBox.warning(self, "Input Error", "Callsign is required for check-in.")
            return

        if self.current_mode == "SKYWARN":
            col_detail = info_val
            col_info = "-"
        else:
            has_traffic = self.chk_traffic.isChecked()
            has_announcement = self.chk_announcement.isChecked()
            col_detail = self.format_additional_attributes(has_traffic, has_announcement)
            col_info = info_val

        for st in self.stations:
            if st["callsign"] == call:
                confirm = QMessageBox.question(
                    self, "Duplicate Station Check-In",
                    f"Station '{call}' is already checked in.\n\nDo you want to update their check-in details and reset status to Active?",
                    QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                    QMessageBox.StandardButton.No
                )
                if confirm == QMessageBox.StandardButton.Yes:
                    st["type"] = st_type
                    st["detail"] = col_detail
                    st["info"] = col_info
                    st["status"] = "Active"
                    st["time_in"] = now_str
                    st["time_out"] = ""
                    self.register_station_callsign(call, col_info if self.current_mode != "SKYWARN" else col_detail)
                    self.refresh_station_table()
                    self.refresh_callsign_combo()
                    self.clear_checkin_inputs()
                    self.station_table.scrollToBottom()
                return

        station_data = {
            "callsign": call,
            "type": st_type,
            "detail": col_detail,
            "info": col_info,
            "status": "Active",
            "time_in": now_str,
            "time_out": ""
        }
        self.stations.append(station_data)
        self.register_station_callsign(call, col_info if self.current_mode != "SKYWARN" else col_detail)
        self.refresh_station_table()
        self.refresh_callsign_combo()
        self.clear_checkin_inputs()
        self.station_table.scrollToBottom()

    def clear_checkin_inputs(self):
        self.input_callsign.setText("")
        if self.input_callsign.completer():
            self.input_callsign.completer().setCompletionPrefix("")
        self.input_roster_loc.clear()
        if self.combo_station_type.count() > 0:
            self.combo_station_type.setCurrentIndex(0)
        self.chk_traffic.setChecked(False)
        self.chk_announcement.setChecked(False)
        self.input_callsign.setFocus()

    def refresh_station_table(self):
        self.station_table.blockSignals(True)
        self.station_table.setRowCount(0)
        total_count = len(self.stations)
        mobile_count = 0

        self.station_table.setUpdatesEnabled(False)
        try:
            for row, st in enumerate(self.stations):
                self.station_table.insertRow(row)

                st_type = st["type"]
                if "Mobile" in st_type:
                    mobile_count += 1

                if self.current_mode == "SKYWARN":
                    items = [
                        QTableWidgetItem(st["callsign"]),
                        QTableWidgetItem(st["type"]),
                        QTableWidgetItem(st.get("detail", "-")),
                        QTableWidgetItem(st["status"]),
                        QTableWidgetItem(st["time_in"]),
                        QTableWidgetItem(st["time_out"])
                    ]
                    status_col_idx = 3
                else:
                    items = [
                        QTableWidgetItem(st["callsign"]),
                        QTableWidgetItem(st["type"]),
                        QTableWidgetItem(st.get("detail", "-")),
                        QTableWidgetItem(st.get("info", "-")),
                        QTableWidgetItem(st["status"]),
                        QTableWidgetItem(st["time_in"]),
                        QTableWidgetItem(st["time_out"])
                    ]
                    status_col_idx = 4

                status = st["status"]
                for col_idx, item in enumerate(items):
                    if status == "Standby":
                        item.setBackground(QColor("#FFFF00"))
                        item.setForeground(QColor("#000000"))
                    elif status == "Checked Out":
                        item.setBackground(QColor("#CC0000"))
                        item.setForeground(QColor("#FFFFFF"))
                    else:
                        item.setBackground(QColor("#FFFFFF"))
                        if col_idx == status_col_idx:
                            item.setForeground(QColor("#008000"))
                        else:
                            item.setForeground(QColor("#000000"))

                    self.station_table.setItem(row, col_idx, item)
        finally:
            self.station_table.setUpdatesEnabled(True)
            self.station_table.blockSignals(False)

        self.lbl_station_counts.setText(f"Total: {total_count}\nMobile: {mobile_count}")

    def refresh_callsign_combo(self):
        current = self.combo_traffic_call.currentText()
        self.combo_traffic_call.clear()
        self.combo_traffic_call.addItem("")
        for st in self.stations:
            if st["status"] != "Checked Out":
                self.combo_traffic_call.addItem(st["callsign"])
        self.combo_traffic_call.setCurrentText(current)

    def toggle_station_standby(self):
        selected = self.station_table.currentRow()
        if selected < 0:
            return
        st = self.stations[selected]
        if st["status"] == "Active":
            st["status"] = "Standby"
        elif st["status"] == "Standby":
            st["status"] = "Active"
        self.refresh_station_table()
        self.station_table.selectRow(selected)

    def checkout_station(self):
        selected = self.station_table.currentRow()
        if selected < 0:
            return
        st = self.stations[selected]
        now_str = datetime.now().strftime("%H:%M")
        st["status"] = "Checked Out"
        st["time_out"] = now_str
        self.refresh_station_table()
        self.refresh_callsign_combo()
        self.station_table.selectRow(selected)

    def log_traffic_item(self):
        call = self.combo_traffic_call.currentText().strip().upper() or "NET"
        traffic_type = self.combo_traffic_type.currentText()
        note = self.input_traffic_note.toPlainText().strip()
        time_val = self.input_traffic_time.text().strip() or datetime.now().strftime("%H:%M")
        loc_val = self.input_traffic_loc.text().strip() or "-"

        if not note:
            QMessageBox.warning(self, "Input Error", "Please provide details for the entry.")
            return

        self.add_activity_log(call, traffic_type, loc_val, note, time_val)
        
        self.combo_traffic_call.setCurrentText("")
        self.input_traffic_note.clear()
        self.input_traffic_loc.clear()
        
        self.combo_traffic_call.setFocus()

    def add_activity_log(self, call, traffic_type, location, note, custom_time=None):
        time_str = custom_time if custom_time else datetime.now().strftime("%H:%M")

        self.activity_table.blockSignals(True)
        try:
            row = self.activity_table.rowCount()
            self.activity_table.insertRow(row)

            if self.current_mode == "SKYWARN":
                activity_item = {
                    "timestamp": time_str,
                    "condition": traffic_type,
                    "location": location,
                    "source": call,
                    "narrative": note
                }
                self.activities.append(activity_item)

                self.activity_table.setItem(row, 0, QTableWidgetItem(time_str))
                self.activity_table.setItem(row, 1, QTableWidgetItem(traffic_type))
                self.activity_table.setItem(row, 2, QTableWidgetItem(location))
                self.activity_table.setItem(row, 3, QTableWidgetItem(call))
                self.activity_table.setItem(row, 4, QTableWidgetItem(note))
            else:
                activity_item = {
                    "timestamp": time_str,
                    "callsign": call,
                    "type": traffic_type,
                    "note": note
                }
                self.activities.append(activity_item)

                self.activity_table.setItem(row, 0, QTableWidgetItem(time_str))
                self.activity_table.setItem(row, 1, QTableWidgetItem(call))

                type_item = QTableWidgetItem(traffic_type)
                if traffic_type == "Emergency":
                    type_item.setForeground(QColor("#FF0000"))
                elif traffic_type == "Priority":
                    type_item.setForeground(QColor("#804000"))

                self.activity_table.setItem(row, 2, type_item)
                self.activity_table.setItem(row, 3, QTableWidgetItem(note))
        finally:
            self.activity_table.blockSignals(False)

        self.activity_table.scrollToBottom()

    def export_xlsx(self):
        if not self.stations:
            QMessageBox.warning(self, "Export XLSX Error", "There are no stations in the roster to export.")
            return

        path, _ = QFileDialog.getSaveFileName(self, "Export XLSX Check-In Roster", "net_checkins.xlsx", "Excel Files (*.xlsx)")
        if not path:
            return

        try:
            export_xlsx_checkin_list(self.stations, path)
            QMessageBox.information(self, "Export Successful", f"Check-in roster successfully exported to XLSX:\n{path}")
        except Exception as e:
            QMessageBox.critical(self, "Export Error", f"Failed to export XLSX file:\n{str(e)}")

    def export_csv(self):
        path, _ = QFileDialog.getSaveFileName(self, "Export CSV Log", "net_log.csv", "CSV Files (*.csv)")
        if not path:
            return

        try:
            with open(path, mode="w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(["=== STATION ROSTER ==="])
                if self.current_mode == "SKYWARN":
                    writer.writerow(["Call", "Station Type", "Location", "Status", "In", "Out"])
                    for st in self.stations:
                        writer.writerow([st.get("callsign", ""), st.get("type", ""), st.get("detail", "-"), st.get("status", ""), st.get("time_in", ""), st.get("time_out", "")])
                else:
                    writer.writerow(["Call", "Station Type", "Additional", "Info / Grid", "Status", "In", "Out"])
                    for st in self.stations:
                        writer.writerow([st.get("callsign", ""), st.get("type", ""), st.get("detail", "-"), st.get("info", "-"), st.get("status", ""), st.get("time_in", ""), st.get("time_out", "")])

                writer.writerow([])
                writer.writerow(["=== NET ACTIVITY & TRAFFIC LOG ==="])
                if self.current_mode == "SKYWARN":
                    writer.writerow(["Time", "Condition", "Location", "Source", "Narrative"])
                    for act in self.activities:
                        writer.writerow([act.get("timestamp", ""), act.get("condition", ""), act.get("location", ""), act.get("source", ""), act.get("narrative", "")])
                else:
                    writer.writerow(["Time", "Callsign", "Traffic Tag", "Details"])
                    for act in self.activities:
                        writer.writerow([act.get("timestamp", ""), act.get("callsign", ""), act.get("type", ""), act.get("note", "")])

            QMessageBox.information(self, "Export Successful", f"Net log successfully exported to:\n{path}")
        except Exception as e:
            QMessageBox.critical(self, "Export Error", f"Failed to export CSV file:\n{str(e)}")

    def export_pdf(self):
        path, _ = QFileDialog.getSaveFileName(self, "Export PDF Summary", "net_report.pdf", "PDF Files (*.pdf)")
        if not path:
            return

        date_str = datetime.now().strftime("%Y-%m-%d %H:%M")

        html = f"""
        <html>
        <head>
            <style>
                body {{ font-family: Tahoma, sans-serif; font-size: 11pt; color: #000; }}
                h1 {{ color: #000080; margin-bottom: 2px; }}
                .meta {{ font-size: 9pt; color: #555; margin-bottom: 20px; }}
                table {{ width: 100%; border-collapse: collapse; margin-top: 10px; margin-bottom: 20px; }}
                th {{ background-color: #000080; color: white; text-align: left; padding: 6px; font-size: 10pt; }}
                td {{ border: 1px solid #808080; padding: 6px; font-size: 9.5pt; }}
                tr.row-standby {{ background-color: #FFFF99; color: #000; }}
                tr.row-checkout {{ background-color: #CC0000; color: #FFF; }}
            </style>
        </head>
        <body>
            <h1>Amateur Radio Net Control Summary</h1>
            <div class="meta">Generated: {date_str} | Profile: {self.current_profile_name} ({self.current_mode} Mode) | NCS: {self.current_ncs_callsign}</div>

            <h3>Station Roster ({len(self.stations)} Stations)</h3>
            <table>
                <tr>
                    <th>Call</th>
                    <th>Station Type</th>
                    <th>{'Location' if self.current_mode == 'SKYWARN' else 'Additional'}</th>
                    {'' if self.current_mode == 'SKYWARN' else '<th>Info / Grid</th>'}
                    <th>Status</th>
                    <th>In</th>
                    <th>Out</th>
                </tr>
        """

        for st in self.stations:
            row_class = ""
            if st.get("status") == "Standby":
                row_class = ' class="row-standby"'
            elif st.get("status") == "Checked Out":
                row_class = ' class="row-checkout"'

            info_td = f"<td>{st.get('info', '-')}</td>" if self.current_mode != 'SKYWARN' else ""

            html += f"""
                <tr{row_class}>
                    <td><b>{st.get('callsign', '')}</b></td>
                    <td>{st.get('type', '')}</td>
                    <td>{st.get('detail', '-')}</td>
                    {info_td}
                    <td>{st.get('status', '')}</td>
                    <td>{st.get('time_in', '')}</td>
                    <td>{st.get('time_out', '')}</td>
                </tr>
            """

        html += f"""
            </table>

            <h3>Net Activity & Traffic Log</h3>
            <table>
                <tr>
                    {"<th>Time</th><th>Condition</th><th>Location</th><th>Source</th><th>Narrative</th>" if self.current_mode == "SKYWARN" else "<th>Time</th><th>Callsign</th><th>Traffic Tag</th><th>Details</th>"}
                </tr>
        """

        for act in self.activities:
            if self.current_mode == "SKYWARN":
                html += f"""
                    <tr>
                        <td>{act.get('timestamp', '')}</td>
                        <td><b>{act.get('condition', '')}</b></td>
                        <td>{act.get('location', '-')}</td>
                        <td>{act.get('source', '')}</td>
                        <td>{act.get('narrative', '')}</td>
                    </tr>
                """
            else:
                html += f"""
                    <tr>
                        <td>{act.get('timestamp', '')}</td>
                        <td><b>{act.get('callsign', '')}</b></td>
                        <td>{act.get('type', '')}</td>
                        <td>{act.get('note', '')}</td>
                    </tr>
                """

        html += """
            </table>
        </body>
        </html>
        """

        try:
            document = QTextDocument()
            document.setHtml(html)

            printer = QPrinter(QPrinter.PrinterMode.HighResolution)
            printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
            printer.setOutputFileName(path)

            document.print(printer)
            QMessageBox.information(self, "Export Successful", f"PDF report successfully saved to:\n{path}")
        except Exception as e:
            QMessageBox.critical(self, "Export Error", f"Failed to export PDF file:\n{str(e)}")
            
    def handle_repeater_badge_click(self):
        prof = self.profiles.get(self.current_profile_name, {})
        rep_cfg = prof.get("repeater", {})

        if not rep_cfg.get("has_backup") or not rep_cfg.get("backup_freq"):
            return

        current_target = "Primary" if self.repeater_badge.is_using_backup else "Backup"
        target_freq = rep_cfg.get("freq" if self.repeater_badge.is_using_backup else "backup_freq")

        confirm = QMessageBox.question(
            self,
            "Switch Repeater",
            f"Switch net operation to the {current_target} repeater ({target_freq})?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No
        )

        if confirm == QMessageBox.StandardButton.Yes:
            new_state = not self.repeater_badge.is_using_backup
            self.repeater_badge.set_repeater_info(rep_cfg, is_using_backup=new_state)

            rep_type = "BACKUP" if new_state else "PRIMARY"
            f_val = rep_cfg.get("backup_freq" if new_state else "freq")
            t_val = rep_cfg.get("backup_tone" if new_state else "tone")

            self.add_activity_log(
                call="NET CONTROL",
                traffic_type="Frequency Change",
                location="-",
                note=f"=== SWITCHED TO {rep_type} REPEATER ({f_val} / {t_val}) ==="
            )
            
    def on_activity_cell_changed(self, item):
        row = item.row()
        col = item.column()
        if row < 0 or row >= len(self.activities):
            return

        self.activity_table.blockSignals(True)
        try:
            new_text = item.text().strip()
            act = self.activities[row]

            if self.current_mode == "SKYWARN":
                if col == 0:
                    act["timestamp"] = new_text
                elif col == 1:
                    act["condition"] = new_text
                    act["type"] = new_text
                elif col == 2:
                    act["location"] = new_text
                elif col == 3:
                    act["source"] = new_text.upper()
                    act["callsign"] = new_text.upper()
                    item.setText(act["source"])
                elif col == 4:
                    act["narrative"] = new_text
                    act["note"] = new_text
            else:
                if col == 0:
                    act["timestamp"] = new_text
                elif col == 1:
                    act["callsign"] = new_text.upper()
                    act["source"] = new_text.upper()
                    item.setText(act["callsign"])
                elif col == 2:
                    act["type"] = new_text
                    act["condition"] = new_text
                elif col == 3:
                    act["note"] = new_text
                    act["narrative"] = new_text
        finally:
            self.activity_table.blockSignals(False)

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = NetControlApp()
    window.show()

    QTimer.singleShot(50, window.prompt_profile_on_startup)

    sys.exit(app.exec())